from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(r"E:\strategy")


def main():
    for path in sorted((ROOT / "lesson").glob("*.pdf")):
        print(
            f"PDF\t{path.name}\tpages={len(PdfReader(str(path)).pages)}"
            f"\tbytes={path.stat().st_size}"
        )

    for path in [
        ROOT / "บริษัท-เถ้าแก่น้อย-ฟู๊ดแอนด์มาร์เก็ตติ้ง-จำกัด-มหาชน.pdf",
        ROOT / "SFAS บริษัท ประชาอาภรณ์ จำกัด.pdf",
    ]:
        print(
            f"PDF\t{path.name}\tpages={len(PdfReader(str(path)).pages)}"
            f"\tbytes={path.stat().st_size}"
        )

    for path in sorted(ROOT.glob("*.docx")):
        doc = Document(str(path))
        print(
            f"DOCX\t{path.name}\tparas={len(doc.paragraphs)}"
            f"\ttables={len(doc.tables)}\tsections={len(doc.sections)}"
            f"\tbytes={path.stat().st_size}"
        )


if __name__ == "__main__":
    main()
