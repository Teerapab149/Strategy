from pathlib import Path
import json

import pdfplumber
from docx import Document


ROOT = Path(r"E:\strategy")
OUT = ROOT / "research" / "curriculum_audit"
OUT.mkdir(parents=True, exist_ok=True)


def extract_pdf(path: Path) -> dict:
    page_lengths = []
    parts = []
    with pdfplumber.open(path) as pdf:
        for idx, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            page_lengths.append(len(text))
            parts.append(f"\n\n===== PAGE {idx} =====\n{text}")
    target = OUT / f"{path.stem}.txt"
    target.write_text("".join(parts), encoding="utf-8")
    return {
        "file": str(path),
        "pages": len(page_lengths),
        "characters": sum(page_lengths),
        "empty_pages": [i + 1 for i, n in enumerate(page_lengths) if n == 0],
        "page_lengths": page_lengths,
        "text_file": str(target),
    }


def extract_docx(path: Path) -> dict:
    doc = Document(str(path))
    blocks = []
    headings = []
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else ""
        blocks.append(f"[P{idx} | {style}] {text}")
        if style.lower().startswith("heading") or style.lower() in {"title", "subtitle"}:
            headings.append({"index": idx, "style": style, "text": text})
    for t_idx, table in enumerate(doc.tables, start=1):
        blocks.append(f"\n===== TABLE {t_idx} ({len(table.rows)}x{len(table.columns)}) =====")
        for row in table.rows:
            blocks.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
    target = OUT / f"{path.stem}_docx.txt"
    target.write_text("\n".join(blocks), encoding="utf-8")
    return {
        "file": str(path),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "headings": headings,
        "text_file": str(target),
    }


def main():
    records = []
    for path in sorted((ROOT / "lesson").glob("*.pdf")):
        records.append(extract_pdf(path))
    for path in [
        ROOT / "บริษัท-เถ้าแก่น้อย-ฟู๊ดแอนด์มาร์เก็ตติ้ง-จำกัด-มหาชน.pdf",
        ROOT / "SFAS บริษัท ประชาอาภรณ์ จำกัด.pdf",
    ]:
        records.append(extract_pdf(path))
    for path in [
        ROOT / "บริษัท ประชาอาภรณ์ จำกัด (มหาชน) เเก้ใหม่.docx",
        ROOT / "research" / "PG_Strategic_Research_Dossier_2026.docx",
    ]:
        records.append(extract_docx(path))
    (OUT / "manifest.json").write_text(
        json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    for record in records:
        if "pages" in record:
            print(
                f"PDF {Path(record['file']).name}: pages={record['pages']} "
                f"chars={record['characters']} empty={record['empty_pages']}"
            )
        else:
            print(
                f"DOCX {Path(record['file']).name}: paras={record['paragraphs']} "
                f"tables={record['tables']} headings={len(record['headings'])}"
            )


if __name__ == "__main__":
    main()
