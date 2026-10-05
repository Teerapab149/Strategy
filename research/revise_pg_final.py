from __future__ import annotations

import math
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(r"E:\strategy")
SOURCE = ROOT / "PG_Strategic_Management_Final_2569.docx"
OUTPUT = ROOT / "PG_Strategic_Management_Final_2569_REVISED.docx"

THAI_FONT = "TH Sarabun New"
BLUE = "2F5597"
LIGHT_BLUE = "D9E2F3"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GRAY = "F2F2F2"


def configure_stdout() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


def set_run_font(run, size=16, bold=None, italic=None, color=None):
    run.font.name = THAI_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), THAI_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), THAI_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), THAI_FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_text(p: Paragraph, text: str, *, size=16, bold=None, italic=None,
                       color=None, style=None, align=None):
    if style:
        # This master exposes built-in style names but its python-docx lookup
        # table resolves the style IDs (Heading1/Heading2/...) reliably.
        try:
            p.style = style
        except KeyError:
            p.style = p.part.styles[style.replace(" ", "")]
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    if align is not None:
        p.alignment = align
    return p


def find_paragraph(doc: Document, starts: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts):
            return p
    raise KeyError(f"Paragraph not found: {starts}")


def insert_paragraph_after(doc: Document, anchor: Paragraph, text: str = "", *, style=None,
                           size=16, bold=None, italic=None, color=None, align=None) -> Paragraph:
    p = doc.add_paragraph()
    anchor._p.addnext(p._p)
    return set_paragraph_text(p, text, size=size, bold=bold, italic=italic,
                              color=color, style=style, align=align)


def insert_paragraph_before(doc: Document, anchor: Paragraph, text: str = "", *, style=None,
                            size=16, bold=None, italic=None, color=None, align=None) -> Paragraph:
    p = anchor.insert_paragraph_before()
    return set_paragraph_text(p, text, size=size, bold=bold, italic=italic,
                              color=color, style=style, align=align)


def remove_paragraph(p: Paragraph):
    parent = p._element.getparent()
    if parent is not None:
        parent.remove(p._element)
    p._p = p._element = None


def set_toc_line(p: Paragraph, title: str, page: int, level: int = 1):
    p.clear()
    p.paragraph_format.left_indent = Inches(0.25 if level == 2 else 0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    r = p.add_run(f"{title}\t{page}")
    set_run_font(r, size=11, bold=(level == 1))
    return p


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, *, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
                  color=None, fill=None):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        set_cell_shading(cell, fill)
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            set_run_font(r, size=size, bold=bold, color=color)


def set_table_borders(table: Table, color="B7B7B7", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def format_table(table: Table, *, widths=None, font_size=10.5, header_size=11,
                 header_fill=LIGHT_BLUE, first_col_left=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for w in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(w))
            grid.append(col)
    for ri, row in enumerate(table.rows):
        if ri == 0:
            set_repeat_header(row)
        for ci, cell in enumerate(row.cells):
            if widths and ci < len(widths):
                set_cell_width(cell, widths[ci])
            fill = header_fill if ri == 0 else None
            align = WD_ALIGN_PARAGRAPH.LEFT if (first_col_left and ci == 0 and ri > 0) else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cell, cell.text, size=header_size if ri == 0 else font_size,
                          bold=(ri == 0), align=align, fill=fill)


def clear_data_rows(table: Table):
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)


def rebuild_table(table: Table, rows, *, section_rows=None, widths=None, font_size=10.5,
                  header_size=11, numeric_cols=None, emphasis_rows=None):
    clear_data_rows(table)
    section_rows = set(section_rows or [])
    numeric_cols = set(numeric_cols or [])
    emphasis_rows = set(emphasis_rows or [])
    header = rows[0]
    for ci, value in enumerate(header):
        set_cell_text(table.rows[0].cells[ci], value, size=header_size, bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER, fill=LIGHT_BLUE)
    for ri, values in enumerate(rows[1:], start=1):
        row = table.add_row()
        if ri in section_rows:
            merged = row.cells[0]
            for c in row.cells[1:]:
                merged = merged.merge(c)
            set_cell_text(row.cells[0], values[0], size=font_size + 0.5, bold=True,
                          align=WD_ALIGN_PARAGRAPH.LEFT, fill=LIGHT_GRAY)
            continue
        for ci, value in enumerate(values):
            fill = LIGHT_ORANGE if ri in emphasis_rows else None
            align = WD_ALIGN_PARAGRAPH.CENTER if ci in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[ci], value, size=font_size,
                          bold=(ri in emphasis_rows and ci in {0, len(values)-1}),
                          align=align, fill=fill)
    format_table(table, widths=widths, font_size=font_size, header_size=header_size)


def add_table_after(doc: Document, anchor: Paragraph, rows, cols, *, widths=None, font_size=9.5):
    table = doc.add_table(rows=rows, cols=cols)
    anchor._p.addnext(table._tbl)
    if widths:
        format_table(table, widths=widths, font_size=font_size, header_size=10.5)
    return table


def xml_paragraph_text(el) -> str:
    return "".join(el.itertext()).strip()


def ensure_source_after(doc: Document, table: Table, text: str):
    nxt = table._tbl.getnext()
    if nxt is not None and nxt.tag == qn("w:p") and xml_paragraph_text(nxt).startswith("ที่มา:"):
        p = Paragraph(nxt, table._parent)
        set_paragraph_text(p, text, size=10, italic=True, color="666666")
        return p
    p = doc.add_paragraph()
    table._tbl.addnext(p._p)
    return set_paragraph_text(p, text, size=10, italic=True, color="666666")


def set_update_fields(doc: Document):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_core_properties(doc: Document):
    doc.core_properties.title = "แผนกลยุทธ์ บริษัท ประชาอาภรณ์ จำกัด (มหาชน) ฉบับปรับปรุง 2569"
    doc.core_properties.subject = "Focused Turnaround + Concentrated Growth, 2570–2574"
    doc.core_properties.author = "คณะผู้จัดทำรายงานวิชา 460-401 การจัดการเชิงกลยุทธ์"
    doc.core_properties.comments = "ปรับปรุงข้อมูลถึง 31 สิงหาคม 2569; Management Target Scenario ไม่ใช่คำแนะนำของบริษัท"


def money(x):
    return f"{x:,.2f}"


def pct(x):
    return f"{x:.2f}"


class StableParagraphDocument:
    """Delegate to python-docx while preserving the master's original paragraph map.

    Newly inserted paragraphs change ``Document.paragraphs`` immediately.  The
    revision plan, however, targets paragraph objects in the untouched master;
    caching that list prevents later insertions from shifting every target.
    Tables and all other document collections remain live through delegation.
    """

    def __init__(self, document):
        self._document = document
        self.paragraphs = list(document.paragraphs)

    def __getattr__(self, name):
        return getattr(self._document, name)


def main():
    configure_stdout()
    doc = StableParagraphDocument(Document(SOURCE))

    # Cover placeholders removed without fabricating names.
    set_paragraph_text(doc.paragraphs[17], "คณะผู้จัดทำรายงานกลุ่ม", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_paragraph_text(doc.paragraphs[23], "อาจารย์ผู้สอนประจำรายวิชา", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Executive summary: revised evidence, reconstructed matrices, disciplined finance language.
    set_paragraph_text(doc.paragraphs[67],
        "บริษัท ประชาอาภรณ์ จำกัด (มหาชน) หรือ PG เป็นผู้ผลิตเครื่องนุ่งห่มและผ้าถักที่ดำเนินธุรกิจมา 45 ปี "
        "มีโรงงานที่ลำพูนและกบินทร์บุรี ผลิตเสื้อผ้าบุรุษและสตรี ชุดกีฬา ชุดว่ายน้ำ ชุดชั้นในชาย "
        "เครื่องแบบองค์กร ชุดทางการแพทย์ และผ้าถัก ทั้งภายใต้แบรนด์ที่ได้รับสิทธิและงาน OEM/ODM "
        "ปี 2568 มียอดขาย 605.20 ล้านบาท ลดลงร้อยละ 21.80 ขาดทุนสุทธิประมาณ 5.56 ล้านบาท "
        "และมีกำไรขั้นต้น 129.52 ล้านบาท [1][2][3]", size=16)
    set_paragraph_text(doc.paragraphs[68],
        "ประเด็นเชิงกลยุทธ์คือธุรกิจหลักยังไม่สร้างผลตอบแทนเพียงพอ กำไรขั้นต้นหักค่าใช้จ่ายขายและบริหารติดลบ "
        "27.97, 24.11 และ 48.56 ล้านบาทในปี 2566–2568 โดยคำว่า “ผลจากธุรกิจหลักโดยประมาณ” ในรายงานนี้หมายถึง "
        "กำไรขั้นต้น − ค่าใช้จ่ายขายและบริหาร มิใช่ EBIT ที่ตรวจสอบแล้ว ปัญหานี้ซ้อนกับการกระจุกตัวของลูกค้า "
        "(I.C.C. ร้อยละ 40.12) สินค้าคงเหลือปลายงวด 337.11 ล้านบาท และสัดส่วนสินทรัพย์ทางการเงินร้อยละ 57.81 "
        "ของสินทรัพย์รวม แบบจำลองบริหารคำนวณระยะเวลาสินค้าคงเหลือปีฐานจากสินค้าคงเหลือเฉลี่ย/ต้นทุนขาย × 365 "
        "ได้ประมาณ 267.75 วัน ซึ่งต่างจากอัตรา 183.70 วันที่บริษัทเปิดเผยเพราะฐานคำนวณต่างกัน [1][2]", size=16)
    set_paragraph_text(doc.paragraphs[69],
        "เมื่อแก้การจำแนก SPI/เครือสหพัฒน์เป็นทรัพยากรภายใน แยก T5 ออกจากช่องว่างเทคโนโลยีภายใน และให้คะแนนใหม่ "
        "EFAS เท่ากับ 2.33 IFAS เท่ากับ 2.64 และ SFAS เท่ากับ 2.48 คะแนนเหล่านี้ไม่ใช่กฎกลไกที่สั่งให้เลือกกลยุทธ์ใด "
        "แต่เมื่ออ่านร่วมกับผลขาดทุนของธุรกิจหลัก Five Forces และช่องว่างคู่แข่ง ยืนยันความจำเป็นของการปรับตำแหน่งเชิงกลยุทธ์ "
        "เมทริกซ์ทางเลือกที่คำนวณใหม่ให้ C = 3.90, B = 2.70 และ A = 2.15 และ C ยังนำทุกกรณีทดสอบความไว", size=16)
    set_paragraph_text(doc.paragraphs[70],
        "กลยุทธ์ที่เลือกยังคงเป็นการฟื้นฟูกิจการแบบมุ่งเน้นควบคู่การเติบโตแบบมุ่งเน้น "
        "(Focused Turnaround + Concentrated Growth) ในเครื่องแบบ ชุดทำงาน ชุดทางการแพทย์ และงาน ODM มูลค่าสูงที่คัดเลือก "
        "โดยใช้การสร้างความแตกต่างแบบมุ่งเน้น (Focused Differentiation) แผนห้ามแข่งขันปริมาณใน commodity OEM "
        "ห้ามขยายกำลังผลิตโดยไม่มีอุปสงค์ผูกพันและกรณีทางธุรกิจ ห้ามเปลี่ยนเป็นแบรนด์ผู้บริโภคเต็มรูปแบบระหว่างแผน "
        "และห้ามรับงานที่ไม่ผ่านเกณฑ์กำไรส่วนเพิ่ม", size=16)
    set_paragraph_text(doc.paragraphs[71],
        "แผน 2570–2574 มีช่วงเตรียมความพร้อมตั้งแต่ปลาย 2569 และใช้ Gate 0–4 ตามลำดับ "
        "วิเคราะห์กำไร พิสูจน์สินค้าและตลาด ขยายเชิงพาณิชย์ ลงทุนอัตโนมัติแบบเลือกเฉพาะ และขยายส่งออกแบบเลือกเฉพาะ "
        "งบประมาณรวมยังอยู่ที่ 102.5–174.0 ล้านบาท โดย P8 เท่ากับ 45–75 ล้านบาทตลอดสามปี หรือเฉลี่ย 15–25 ล้านบาทต่อปี "
        "ผลเป้าหมายปี 2574 คือยอดขายประมาณ 810 ล้านบาท อัตรากำไรขั้นต้นร้อยละ 27.5 ค่าใช้จ่ายขายและบริหารร้อยละ 20.0 "
        "ผลจากธุรกิจหลักประมาณร้อยละ 7.5 และระยะเวลาสินค้าคงเหลือตามสูตรแบบจำลองประมาณ 110 วัน "
        "อัตราเติบโตเฉลี่ยโดยนัยร้อยละ 4.98 อยู่ในช่วงเดียวกับกรอบการเติบโตร้อยละ 4–5 ต่อปีที่บริษัทเคยประกาศ "
        "แต่ไม่ใช่การอ้างว่าเป้าหมายปี 2574 เป็นเป้าหมายที่บริษัทประกาศโดยตรง", size=16)
    set_paragraph_text(doc.paragraphs[72],
        "ข้อจำกัดสำคัญคือไม่พบข้อมูลสาธารณะเพียงพอเกี่ยวกับกำไรแยกตามกลุ่มผลิตภัณฑ์ ลูกค้า และช่องทาง "
        "จึงกำหนด P1 ให้พิสูจน์กำไรครอบคลุมยอดขายอย่างน้อยร้อยละ 95 ภายใน 90 วันเป็น Gate 0 "
        "ก่อนอนุมัติการขยายเชิงพาณิชย์หรือ P8 ฉากทัศน์งบการเงินในบทที่ 6 เป็น “Management Target Scenario – not company guidance” "
        "และทุกยอดที่ไม่มีหลักฐานสาธารณะระบุเป็นสมมติฐานเชิงวางแผนของผู้จัดทำ", size=16)

    # Current board and shareholder facts.
    set_paragraph_text(doc.paragraphs[80], "2.3 คณะกรรมการบริษัทปัจจุบัน", style="Heading 2", size=16, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[81],
        "ตามมติคณะกรรมการบริษัทวันที่ 12 พฤษภาคม 2569 คณะกรรมการบริษัทปัจจุบันมี 10 คน และเป็นกรรมการอิสระ 4 คน "
        "(ร้อยละ 40) ได้แก่ (1) Chailert Manoonpol – Chairman (2) Sunan Niyomnaitham (3) Somporn Tiyaviboonsiri "
        "(4) Dusadee Soontrontumrong (5) Duangrudee Milintanggul (6) Suthida Jongjenkit "
        "(7) General Konecharnart Chunnabhata – Independent Director (8) Nuchanart Thammanomai – Independent Director "
        "(9) Vittawat Panpanich – Independent Director และ (10) Thanapat Wongwaitanasakul – Independent Director [22]", size=16)
    set_paragraph_text(doc.paragraphs[82],
        "การเปลี่ยนแปลงกรรมการในปี 2568 ภายหลัง SPI เข้าถือหุ้นเป็นเหตุการณ์ทางธรรมาภิบาลในอดีตและยังใช้ตีความทิศทางได้ "
        "แต่ไม่ควรนำรายชื่อหรือจำนวนกรรมการปี 2568 มาแทนสถานะปัจจุบัน ณ 12 พฤษภาคม 2569 "
        "เชิงกลยุทธ์ คณะกรรมการชุดปัจจุบันควรกำกับ Gate 0–4 การจัดสรรทุน และความเสี่ยงจากรายการหรือช่องทางในเครือโดยแยกบทบาทให้ชัดเจน", size=16)
    set_paragraph_text(doc.paragraphs[76],
        "การเปลี่ยนแปลงสำคัญคือโครงสร้างผู้ถือหุ้น โดย SPI เป็นผู้ถือหุ้นใหญ่และถือหุ้น 45,117,830 หุ้น "
        "คิดเป็นร้อยละ 46.998 ตามทะเบียนผู้ถือหุ้น ณ 13 พฤษภาคม 2569 [23] ความสัมพันธ์กับ SPI และเครือสหพัฒน์ "
        "เป็นความสัมพันธ์ด้านเจ้าของ การกำกับดูแล ทรัพยากร และเครือข่ายขององค์กร จึงประเมินเป็นปัจจัยภายในและ Strategic Enabler "
        "ไม่ใช่โอกาสภายนอกใน EFAS", size=16)
    set_paragraph_text(doc.paragraphs[78],
        "โครงสร้างการกำกับดูแลที่เปิดเผยใน One Report ปี 2568 ใช้คณะกรรมการบริษัทเป็นองค์กรสูงสุดและมีคณะกรรมการชุดย่อย "
        "สำหรับงานตรวจสอบ สรรหาและค่าตอบแทน ความเสี่ยง กำกับดูแลกิจการ/ความยั่งยืน และบริหาร [1] "
        "รายชื่อกรรมการปัจจุบันอัปเดตตามมติวันที่ 12 พฤษภาคม 2569 ในหัวข้อ 2.3 [22] ส่วนสายปฏิบัติการจัดตามหน้าที่และหน่วยผลิต "
        "ได้แก่ สำนักงานใหญ่ โรงงานลำพูน และโรงงานกบินทร์บุรี", size=16)
    set_paragraph_text(doc.paragraphs[79],
        "ข้อสังเกตเชิงกลยุทธ์: จากผังองค์กรและข้อมูลสาธารณะที่ตรวจสอบ ไม่พบหลักฐานเพียงพอว่ามีหน่วยธุรกิจที่รับผิดชอบกำไรขาดทุน "
        "แยกตามกลุ่มลูกค้าหรือผลิตภัณฑ์ เมื่ออ่านร่วมกับผลขาดทุนของธุรกิจหลัก จึงเสนอให้ P1 ยืนยันโครงสร้างผู้รับผิดชอบและจัดตั้ง "
        "accountable P&L ตามสนามที่เลือกในบทที่ 5 ข้อความนี้เป็นข้ออนุมานจากข้อมูลสาธารณะ ไม่ใช่ข้อเท็จจริงภายในที่บริษัทรับรอง", size=16)
    set_paragraph_text(doc.paragraphs[89],
        "การประเมิน: เป้าหมายระยะยาวที่เปิดเผยเป็นเป้าหมายด้านรายได้ และจากเอกสารสาธารณะที่ตรวจสอบไม่พบเป้าหมายเชิงปริมาณ "
        "สำหรับผลของธุรกิจหลักกำกับ ทั้งที่ผลจากธุรกิจหลักโดยประมาณติดลบ แนวทางทั้งห้าข้อเป็นกิจกรรมที่ยังไม่ระบุชัดว่าจะชนะหรือเลิกแข่งในสนามใด "
        "ปี 2568 ยอดขาย PG ลดร้อยละ 21.80 และ core proxy ติดลบมากที่สุดในรอบสามปี ขณะที่การส่งออกเสื้อผ้าสำเร็จรูปไทยขยายตัวร้อยละ 5.61 "
        "จึงมีหลักฐานว่าผลถดถอยไม่ได้มาจากภาวะอุตสาหกรรมเพียงอย่างเดียว [1][3][8]", size=16)

    # Historical financial table: retain disclosed company ratio and make modelling denominator explicit.
    t0 = doc.tables[0]
    set_cell_text(t0.cell(16, 0), "ระยะเวลาขายสินค้าคงเหลือ/วงจรเงินสดที่บริษัทเปิดเผย (วัน)", size=10.5)
    set_cell_text(t0.cell(16, 4), "ข้อเท็จจริง (วิธีคำนวณของบริษัท)", size=10.5)
    set_paragraph_text(doc.paragraphs[93],
        "ที่มา: แบบ 56-1 One Report ปี 2568 [1]; งบการเงินตรวจสอบแล้วปี 2568 [2]; คำชี้แจงผลการดำเนินงานปี 2568 [3] "
        "หมายเหตุ: ผลจากธุรกิจหลักโดยประมาณ = กำไรขั้นต้น − ค่าใช้จ่ายขายและบริหาร มิใช่ EBIT ที่ตรวจสอบแล้ว "
        "ส่วนระยะเวลาสินค้าคงเหลือ 183.70 วันเป็นอัตราที่บริษัทเปิดเผย แบบจำลองในบทที่ 6 ใช้สินค้าคงเหลือเฉลี่ย/ต้นทุนขาย × 365", size=10, italic=True, color="666666")

    # Latest H1/2569 gets its own section and corrected causality.
    p103 = doc.paragraphs[103]
    set_paragraph_text(p103, "2.7 ผลการดำเนินงานล่าสุด: งวดหกเดือนแรกปี 2569", style="Heading 2", size=16, bold=True, color=BLUE)
    insert_paragraph_after(doc, p103, "ตารางที่ 3 ผลการดำเนินงานงวดหกเดือนแรก ปี 2569 เทียบกับปี 2568 (หน่วย: ล้านบาท)",
                           size=14, bold=True)
    set_paragraph_text(doc.paragraphs[105],
        "นัยเชิงกลยุทธ์ต่อ PG: ครึ่งแรกปี 2569 มียอดขาย 285.05 ล้านบาท กำไรขั้นต้น 59.35 ล้านบาท "
        "ค่าใช้จ่ายขายและบริหาร 76.21 ล้านบาท และผลจากธุรกิจหลักโดยประมาณ -16.86 ล้านบาท "
        "กำไรสุทธิ 17.92 ล้านบาทดีขึ้นเมื่อเทียบปีก่อน โดยแรงขับหลักคือการลดค่าใช้จ่ายพนักงานใน SG&A 15.88 ล้านบาท "
        "และการพลิกจากขาดทุนตีมูลค่าเงินลงทุน 22.67 ล้านบาทเป็นกำไร 7.36 ล้านบาท รายได้อื่น 29.77 ล้านบาทยังมีส่วนสำคัญต่อกำไรที่รายงาน "
        "แต่ลดลงร้อยละ 8.70 จึงไม่ใช่สาเหตุของการดีขึ้น YoY ข้อสรุปคือกำไรสุทธิฟื้น แต่ธุรกิจหลักยังไม่คุ้มค่าใช้จ่าย [4]", size=16)

    # External analysis: exact export value, cautious DPP wording, corrected factor boundary.
    t3 = doc.tables[3]
    set_cell_text(t3.cell(6, 1),
        "สำนักงานเศรษฐกิจอุตสาหกรรมรายงานว่าการส่งออกเสื้อผ้าสำเร็จรูปไทยปี 2568 มีมูลค่าประมาณ 2,018.77 ล้านดอลลาร์สหรัฐ "
        "ขยายตัวร้อยละ 5.61 โดยหมวดที่เติบโตรวมถึงชุดกีฬา เสื้อโปโล เสื้อผ้าเด็กอ่อน และชุดชั้นใน [8]", size=10)
    set_cell_text(t3.cell(8, 1),
        "ยุทธศาสตร์สิ่งทอของ EU และ ESPR กำหนดทิศทาง/กรอบข้อมูลความยั่งยืนระดับผลิตภัณฑ์และ Digital Product Passport "
        "แต่รายละเอียดและกำหนดเวลาบังคับสำหรับสิ่งทอต้องเป็นไปตามกฎเฉพาะผลิตภัณฑ์ที่ใช้บังคับ จึงไม่สมมติวันบังคับสากลสำหรับเสื้อผ้าทุกชนิด [10][11]", size=10)
    set_cell_text(t3.cell(10, 1),
        "Regulation (EU) 2024/3015 ว่าด้วยการห้ามผลิตภัณฑ์ที่ผลิตด้วยแรงงานบังคับใช้กับสินค้าที่วางในตลาด EU ตั้งแต่ 14 ธันวาคม 2570 "
        "โดยสิ่งทอเป็นหนึ่งในอุตสาหกรรมเสี่ยงที่ทางการระบุ [12]", size=10)
    set_paragraph_text(doc.paragraphs[113],
        "ข้อแก้ไขสำคัญ: เครื่องนุ่งห่มไม่ใช่ภาคที่อยู่ในขอบเขต CBAM โดยตรง ณ วันที่ตัดข้อมูล "
        "ภาคที่ครอบคลุมคือซีเมนต์ เหล็กและเหล็กกล้า อะลูมิเนียม ปุ๋ย ไฟฟ้า และไฮโดรเจน [18] "
        "แรงกดดันที่เกี่ยวกับ PG โดยตรงกว่าคือ circularity, ecodesign, product information และ forced labour "
        "ทั้งนี้ EU Textile Strategy และ ESPR เป็นกรอบทิศทางของ DPP; รายงานไม่กำหนดวันบังคับสำหรับเสื้อผ้าทุกชนิดหากยังไม่มีกฎเฉพาะผลิตภัณฑ์รองรับ", size=16)
    set_cell_text(doc.tables[6].cell(16, 6),
        "จากข้อมูลสาธารณะที่ตรวจสอบ ไม่พบหลักฐานเพียงพอให้ยืนยันขอบเขตระบบเติมสินค้าตามสัญญาของแต่ละบริษัท "
        "คะแนนจึงเป็นประมาณการของผู้จัดทำ ไม่ใช่ข้อเท็จจริงขององค์กร", size=9.5)
    set_paragraph_text(doc.paragraphs[126],
        "นัยเชิงกลยุทธ์ต่อ PG: ในสนาม OEM ส่งออก PG ตามหลัง 2.20 คะแนน โดยช่องว่างหลักอยู่ที่ขนาดและฐานผลิตหลายประเทศ "
        "แต่ในสนามเครื่องแบบและงานโครงการ PG ตามหลัง 0.95 คะแนน และมีฐานผ้าถัก การควบคุมคุณภาพ และผลิตภัณฑ์เชิงฟังก์ชันที่ใช้ต่อยอดได้ "
        "Warrix จึงใช้เป็น benchmark ใกล้เคียงด้านงานโครงการ/เครื่องแบบ มิใช่คู่แข่งตรงสมบูรณ์ ส่วน Sabina และ MC Group ใช้เพียงเทียบเศรษฐศาสตร์ของแบรนด์/ช่องทาง "
        "ไม่ใช้สมมติว่า PG จะได้อัตรากำไรขั้นต้นเท่ากัน คะแนนคู่แข่งทั้งหมดเป็น analyst estimates จากข้อมูลสาธารณะ", size=16)

    set_paragraph_text(doc.paragraphs[129],
        "โอกาสประกอบด้วย O1 อุปสงค์เครื่องแบบและชุดทำงานเชิงฟังก์ชันในกลุ่มการแพทย์ บริการ และอุตสาหกรรม "
        "O2 กฎเกณฑ์และเกณฑ์จัดซื้อสากลที่ให้ค่ากับการตรวจสอบย้อนกลับและเศรษฐกิจหมุนเวียน "
        "O3 การกลับมาขยายตัวของการส่งออกเครื่องนุ่งห่มไทยในบางหมวด และ O4 ต้นทุนการเข้าถึงเครื่องมือดิจิทัลสำหรับงาน B2B ที่ลดลง "
        "ไม่รวม SPI/เครือสหพัฒน์ซึ่งเป็นปัจจัยภายใน", size=16)
    set_paragraph_text(doc.paragraphs[130],
        "อุปสรรคประกอบด้วย T1 การแข่งขันด้านต้นทุนจากผู้ผลิตต่างประเทศและสินค้านำเข้าราคาต่ำ "
        "T2 อำนาจต่อรองของผู้ซื้อในอุตสาหกรรมรับจ้างผลิตและค้าปลีก T3 กำลังซื้อในประเทศอ่อนแอและการเปลี่ยนนโยบายจัดซื้อของลูกค้ารายใหญ่ "
        "T4 ความไม่แน่นอนจากมาตรการภาษีการค้าและอัตราแลกเปลี่ยน และ T5 มาตรฐานอุตสาหกรรมด้าน automation, response time "
        "และ digital manufacturing ที่สูงขึ้น ส่วนช่องว่างเทคโนโลยี/ผลิตภาพเฉพาะของ PG จัดเป็น W4", size=16)

    efas_rows = [
        ["ปัจจัยภายนอก (External Factors)", "น้ำหนัก", "Rating", "คะแนน", "เหตุผลของน้ำหนักและคะแนน"],
        ["โอกาส (Opportunities)", "", "", "", ""],
        ["O1 อุปสงค์เครื่องแบบและชุดทำงานเชิงฟังก์ชันในกลุ่มการแพทย์ บริการ และอุตสาหกรรม", "0.14", "3", "0.42",
         "สำคัญสูงสุดในกลุ่มโอกาสและตรงกับฐานผลิตภัณฑ์ของ PG; Rating 3 เพราะมีชุดทางการแพทย์และลูกค้าองค์กรแล้ว แต่ยังไม่เป็นหน่วยธุรกิจที่วัดกำไรแยก"],
        ["O2 กฎและเกณฑ์จัดซื้อสากลด้าน traceability/circularity", "0.10", "3", "0.30",
         "กำหนดการเข้าถึงลูกค้ามูลค่าสูง; Rating 3 เพราะมี ISO 14001/GRS แต่ยังไม่มีข้อมูลระดับ SKU/ล็อตครบ"],
        ["O3 การส่งออกเสื้อผ้าสำเร็จรูปไทยบางหมวดกลับมาขยายตัว", "0.07", "2", "0.14",
         "OIE รายงาน 2,018.77 ล้านดอลลาร์สหรัฐ +5.61%; Rating 2 เพราะยอดส่งออก PG กลับลด 34.36%"],
        ["O4 ต้นทุนเข้าถึงเครื่องมือดิจิทัล B2B ลดลง", "0.07", "2", "0.14",
         "ช่วยผู้ผลิตขนาดกลางยกระดับ CRM/costing/S&OP; Rating 2 เพราะ PG ยังไม่เปิดเผยระบบเชิงพาณิชย์ดังกล่าว"],
        ["อุปสรรค (Threats)", "", "", "", ""],
        ["T1 การแข่งขันด้านต้นทุนจากผู้ผลิตต่างประเทศและสินค้านำเข้าราคาต่ำ", "0.16", "2", "0.32",
         "กระทบเพดานอัตรากำไรโดยตรง; Rating 2 เพราะการตอบสนองยังไม่เปลี่ยนสนามการแข่งขันอย่างเป็นระบบ"],
        ["T2 อำนาจต่อรองของผู้ซื้อในอุตสาหกรรมรับจ้างผลิตและค้าปลีก", "0.14", "2", "0.28",
         "แรงกดดันสูงมาก; Rating 2 เพราะยังไม่มีหลักฐานสัญญาระยะกลางหรือกลไกปรับราคาที่ลดอำนาจผู้ซื้อ"],
        ["T3 กำลังซื้อในประเทศอ่อนแอและการเปลี่ยนนโยบายจัดซื้อของลูกค้ารายใหญ่", "0.12", "2", "0.24",
         "กระทบรายได้ในประเทศส่วนใหญ่; Rating 2 เพราะยอดขายในประเทศและครึ่งแรก 2569 ยังลดลง"],
        ["T4 ความไม่แน่นอนจากภาษีการค้าและอัตราแลกเปลี่ยน", "0.09", "3", "0.27",
         "กระทบส่วนส่งออก; Rating 3 เพราะมี natural hedge บัญชีเงินตราต่างประเทศ และ forward ตามนโยบายบริษัท"],
        ["T5 benchmark อุตสาหกรรมด้าน automation, response time และ digital manufacturing สูงขึ้น", "0.11", "2", "0.22",
         "กำหนดความสามารถแข่งขันระยะกลาง; Rating 2 เพราะ PG เริ่มกึ่งอัตโนมัติ แต่ยังไม่เปิดเผยผลลัพธ์ผลิตภาพ/ความเร็ว"],
        ["รวม", "1.00", "", "2.33", ""],
    ]
    rebuild_table(doc.tables[7], efas_rows, section_rows={1, 6}, widths=[3150, 650, 650, 700, 3870],
                  font_size=9.5, numeric_cols={1, 2, 3}, emphasis_rows={12})
    set_paragraph_text(doc.paragraphs[135],
        "นัยเชิงกลยุทธ์ต่อ PG: EFAS รวม 2.33 สะท้อนว่าการตอบสนองต่อโอกาสและอุปสรรคสำคัญยังต่ำกว่าระดับปานกลาง "
        "แต่คะแนนต่ำกว่า 3.00 ไม่ใช่กฎอัตโนมัติให้เลือกกลยุทธ์ใด เมื่ออ่านร่วมกับผลธุรกิจหลัก Five Forces และการเทียบคู่แข่ง "
        "หลักฐานสนับสนุนให้ PG ย้ายจากการแข่งขันด้านปริมาณ/ราคาไปสู่สนามที่ O1/O2 เชื่อมกับความสามารถภายในได้", size=16)

    # VRIO uses one consistent scale and avoids categorical overclaim.
    vrio = doc.tables[8]
    vrio_values = [
        ("ใช่", "บางส่วน", "ไม่ใช่", "บางส่วน"),
        ("ใช่", "บางส่วน", "บางส่วน", "บางส่วน"),
        ("ใช่", "ไม่ใช่", "ไม่ใช่", "ใช่"),
        ("ใช่", "บางส่วน", "บางส่วน", "บางส่วน"),
        ("ใช่", "ใช่", "บางส่วน", "บางส่วน"),
        ("ใช่", "บางส่วน", "บางส่วน", "บางส่วน"),
        ("ใช่", "บางส่วน", "บางส่วน", "ใช่"),
        ("ไม่ใช่", "ไม่ใช่", "ไม่ใช่", "ไม่ใช่"),
    ]
    for ri, vals in enumerate(vrio_values, start=1):
        for ci, val in enumerate(vals, start=1):
            set_cell_text(vrio.cell(ri, ci), val, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(vrio.cell(5, 0), "ความสัมพันธ์กับ SPI, I.C.C. และเครือสหพัฒน์", size=9.5)
    set_paragraph_text(doc.paragraphs[141],
        "ความสามารถที่มีหลักฐานสาธารณะสนับสนุนมีสี่กลุ่ม คือ (1) การบูรณาการตั้งแต่ผ้าถักถึงเสื้อผ้าสำเร็จรูป "
        "(2) การพัฒนาผลิตภัณฑ์เชิงฟังก์ชัน (3) ประสบการณ์รับงานเครื่องแบบองค์กรและชุดทางการแพทย์ และ (4) การบริหารคู่ค้า/การจัดหา "
        "ส่วนการวิเคราะห์กำไรรายลูกค้าและ SKU, account-based B2B, บริการเติมสินค้าตามสัญญา และ traceability ระดับล็อต "
        "เป็นด้านที่ไม่พบข้อมูลสาธารณะเพียงพอให้ยืนยันความพร้อม จึงแปลงเป็นสมมติฐานความสามารถที่ต้องพิสูจน์ผ่าน P1–P7 ในบทที่ 5", size=16)
    set_paragraph_text(doc.paragraphs[145],
        "ความสามารถที่มีหลักฐานสนับสนุน V/R/I มากที่สุดเมื่อเทียบกับความสามารถอื่นของ PG คือการบูรณาการผ้าถัก "
        "คุณสมบัติเชิงฟังก์ชัน การออกแบบ การควบคุมคุณภาพ และการตัดเย็บในองค์กรเดียว แต่ผลประเมินยังเป็น “บางส่วน” ใน R/I/O "
        "เพราะข้อมูลสาธารณะไม่ยืนยันความหายากแบบเด็ดขาดและยังไม่เห็นโครงสร้างที่แปลงความสามารถเป็นอัตรากำไร", size=16)
    set_paragraph_text(doc.paragraphs[146],
        "นัยเชิงกลยุทธ์ต่อ PG: ไม่ควรอ้างว่ามีความได้เปรียบยั่งยืนแล้ว สิ่งที่ควรสร้างคือความสามารถในการออกแบบและส่งมอบ "
        "โซลูชันเครื่องแบบตามภารกิจ พร้อมหลักฐานคุณสมบัติ/ความยั่งยืนระดับ SKU และล็อต และบริการเติมสินค้าตามสัญญา "
        "จากข้อมูลสาธารณะที่ตรวจสอบ ไม่พบหลักฐานเพียงพอให้ยืนยันว่าคู่แข่งรายใดมีหรือไม่มีระบบดังกล่าวครบถ้วน "
        "จึงใช้เป็นสมมติฐานการแข่งขันที่ต้องทดสอบใน Gate 1 ไม่ใช่ข้อเท็จจริงเชิงลบเกี่ยวกับคู่แข่ง", size=16)

    set_paragraph_text(doc.paragraphs[151],
        "จุดแข็งประกอบด้วย S1 ฐานะการเงินแข็งแรงและแทบไม่มีภาระหนี้ S2 ความสามารถบูรณาการผ้าถัก–ออกแบบ–แพตเทิร์น–ตัดเย็บ "
        "S3 ผลิตภัณฑ์เชิงฟังก์ชันและมาตรฐานที่ได้รับการรับรอง S4 ประสบการณ์/ฐานลูกค้าเครื่องแบบและชุดทางการแพทย์ "
        "S5 ฐานคู่ค้าที่กระจายตัวและระบบบริหารคู่ค้า และ S6 ความสัมพันธ์กับ SPI/I.C.C./เครือสหพัฒน์ในฐานะ Strategic Enabler ภายใน", size=16)
    set_paragraph_text(doc.paragraphs[152],
        "จุดอ่อนประกอบด้วย W1 ธุรกิจหลักขาดทุนและ SG&A สูงกว่ากำไรขั้นต้น W2 ลูกค้า/ช่องทางกระจุกตัว "
        "W3 สินค้าคงเหลือและเงินทุนหมุนเวียน W4 ช่องว่างด้านเทคโนโลยี ผลิตภาพ ความเร็วตอบสนอง และข้อมูลปฏิบัติการเมื่อเทียบ benchmark อุตสาหกรรม "
        "W5 การจัดสรรทุนไม่เชื่อมกับผลตอบแทนธุรกิจหลัก และ W6 ขาดโครงสร้าง/ข้อมูลสำหรับกำไรรายกลุ่มลูกค้า", size=16)

    ifas_rows = [
        ["ปัจจัยภายใน (Internal Factors)", "น้ำหนัก", "Rating", "คะแนน", "เหตุผลของน้ำหนักและคะแนน"],
        ["จุดแข็ง (Strengths)", "", "", "", ""],
        ["S1 ฐานะการเงินแข็งแรงและแทบไม่มีภาระหนี้", "0.09", "4", "0.36", "เป็นแหล่งความยืดหยุ่นด้านทุน; Rating 4 เพราะสภาพคล่องและ D/E แข็งแรง แต่การจัดสรรทุนยังต้องปรับ"],
        ["S2 บูรณาการผ้าถัก–ออกแบบ–แพตเทิร์น–ตัดเย็บ", "0.11", "4", "0.44", "ฐานสำคัญของความแตกต่างในสนามที่เลือกและใช้งานจริง แต่ยังไม่แปลงเป็นอัตรากำไรเต็มที่"],
        ["S3 ผลิตภัณฑ์เชิงฟังก์ชันและมาตรฐานภายนอก", "0.09", "4", "0.36", "มีผลิตภัณฑ์/ใบรับรองจริงและเป็น entry criterion ของลูกค้าองค์กร/ส่งออก"],
        ["S4 ประสบการณ์และฐานลูกค้าเครื่องแบบ/ชุดทางการแพทย์", "0.07", "3", "0.21", "ลดเวลาสู่รายได้ แต่ยังไม่มี P&L และทีมขายเฉพาะ"],
        ["S5 ฐานคู่ค้ากระจายตัวและระบบบริหารคู่ค้า", "0.04", "4", "0.16", "222 ราย ไม่มีรายใดเกิน 10% และมีระบบประเมินคู่ค้า"],
        ["S6 ความสัมพันธ์กับ SPI, I.C.C. และเครือสหพัฒน์", "0.07", "3", "0.21", "เป็น Strategic Enabler ด้าน governance/resource/network; Rating 3 เพราะยังไม่มีผลเชิงพาณิชย์ใหม่ที่วัดได้และมี concentration risk"],
        ["จุดอ่อน (Weaknesses)", "", "", "", ""],
        ["W1 ธุรกิจหลักขาดทุนต่อเนื่องและ SG&A สูงกว่ากำไรขั้นต้น", "0.16", "1", "0.16", "กำหนดความอยู่รอด; Rating 1 เพราะ H1/2569 ยังมี core proxy -16.86 ล้านบาท"],
        ["W2 ลูกค้าและช่องทางกระจุกตัว", "0.11", "2", "0.22", "I.C.C. 40.12% ของยอดขายรวมและผลกระทบเกิดจริง"],
        ["W3 สินค้าคงเหลือและเงินทุนหมุนเวียน", "0.09", "2", "0.18", "สินค้าคงเหลือ 337.11 ล้านบาท; DIO แบบจำลอง 267.75 วัน"],
        ["W4 ช่องว่างเทคโนโลยี ผลิตภาพ response time และข้อมูลปฏิบัติการ", "0.06", "2", "0.12", "เป็นช่องว่างเฉพาะ PG เมื่อเทียบ benchmark; ยังไม่มี KPI ผลิตภาพ/ส่งมอบสาธารณะ"],
        ["W5 การจัดสรรทุนไม่เชื่อมผลตอบแทนธุรกิจหลัก", "0.06", "2", "0.12", "สินทรัพย์ทางการเงินสูง ขณะที่ธุรกิจหลักขาดทุนและมีการจ่ายปันผล"],
        ["W6 ขาดโครงสร้างและข้อมูลกำไรรายกลุ่มลูกค้า", "0.05", "2", "0.10", "ไม่พบหน่วยธุรกิจ P&L หรือกำไรแยกกลุ่มในข้อมูลสาธารณะ"],
        ["รวม", "1.00", "", "2.64", ""],
    ]
    rebuild_table(doc.tables[10], ifas_rows, section_rows={1, 8}, widths=[3150, 650, 650, 700, 3870],
                  font_size=9.5, numeric_cols={1, 2, 3}, emphasis_rows={15})
    set_paragraph_text(doc.paragraphs[155],
        "นัยเชิงกลยุทธ์ต่อ PG: IFAS รวม 2.64 แสดงว่าทรัพยากรที่มีอยู่ยังไม่ถูกจัดองค์กรให้เกิดผลลัพธ์ "
        "S2/S3 เป็นฐานความแตกต่าง ส่วน S6 เป็น enabler ไม่ใช่ external opportunity จุดอ่อนเร่งด่วนคือ W1 "
        "และต้องแก้ W4 ด้วยการวัดผลก่อนลงทุน ไม่ใช่สรุปว่าต้องซื้อเครื่องจักรเพราะ benchmark ภายนอกสูงขึ้น", size=16)

    # SFAS reconstructed after corrected classification.
    sf_rows = [
        ["ปัจจัยเชิงกลยุทธ์", "น้ำหนัก", "Rating", "คะแนน", "สั้น", "กลาง", "ยาว", "เหตุผลและผลเชิงกลยุทธ์"],
        ["S1 ฐานะการเงินแข็งแรง", "0.04", "4", "0.16", "●", "●", "", "ทำให้ใช้ stage gate และไม่ต้องก่อหนี้ภายใต้ฉากทัศน์"],
        ["S2 บูรณาการผ้าถัก–ออกแบบ–แพตเทิร์น–ตัดเย็บ", "0.09", "4", "0.36", "", "●", "●", "แกนความแตกต่างของงานเชิงฟังก์ชัน"],
        ["S3 ผลิตภัณฑ์เชิงฟังก์ชันและมาตรฐาน", "0.07", "4", "0.28", "", "●", "●", "แปลงเป็นข้อกำหนด/ผลทดสอบราย SKU"],
        ["S4 ฐานลูกค้าเครื่องแบบและชุดทางการแพทย์", "0.04", "3", "0.12", "●", "●", "", "ฐาน pilot/reference customer"],
        ["S6 เครือข่าย SPI/I.C.C./เครือสหพัฒน์", "0.05", "3", "0.15", "●", "●", "●", "Strategic Enabler ภายในและต้องมีกติกาช่องทาง"],
        ["W1 ธุรกิจหลักขาดทุน", "0.17", "1", "0.17", "●", "", "", "Gate 0 ต้องมาก่อนการเติบโตและ capex"],
        ["W2 ลูกค้า/ช่องทางกระจุกตัว", "0.08", "2", "0.16", "", "●", "", "ต้องเพิ่ม direct B2B และสัญญาระยะกลาง"],
        ["W3 สินค้าคงเหลือและเงินทุนหมุนเวียน", "0.07", "2", "0.14", "●", "●", "", "ใช้ S&OP/MTO/contract replenishment"],
        ["W4 ช่องว่างเทคโนโลยี ผลิตภาพ และข้อมูล", "0.07", "2", "0.14", "●", "●", "●", "ต้องวัด baseline/pilot ก่อน P8"],
        ["O1 อุปสงค์ functional uniform/workwear", "0.11", "3", "0.33", "", "●", "●", "ตลาดที่ตรงกับ S2/S3/S4"],
        ["O2 เกณฑ์ traceability/circularity", "0.05", "3", "0.15", "●", "●", "●", "เร่ง minimum viable compliance ก่อน 14 ธ.ค. 2570"],
        ["T1 การแข่งขันด้านต้นทุน", "0.07", "2", "0.14", "●", "●", "●", "ห้ามแข่ง commodity OEM ด้วยปริมาณ"],
        ["T2 อำนาจต่อรองของผู้ซื้อ", "0.04", "2", "0.08", "●", "●", "", "ใช้ profitability hurdle และ contract terms"],
        ["T5 benchmark automation/response time/digital manufacturing สูงขึ้น", "0.05", "2", "0.10", "", "●", "●", "ตอบด้วย selective automation หลัง Gate 2"],
        ["รวม", "1.00", "", "2.48", "", "", "", ""],
    ]
    rebuild_table(doc.tables[11], sf_rows, widths=[2600, 570, 570, 620, 430, 430, 430, 3370],
                  font_size=9.0, numeric_cols={1, 2, 3, 4, 5, 6}, emphasis_rows={15})
    set_paragraph_text(doc.paragraphs[158],
        "SFAS คัด 14 ปัจจัยจาก EFAS/IFAS หลังแก้การจำแนก โดยให้น้ำหนักใหม่รวม 1.00 และคงเฉพาะปัจจัยที่กำหนดสนาม "
        "ความอยู่รอด การนำไปใช้ หรือความเสี่ยงข้ามเวลา ปัจจัย SPI/เครือสหพัฒน์ปรากฏเป็น S6 ภายใน ไม่ปรากฏซ้ำเป็นโอกาสภายนอก", size=16)
    set_paragraph_text(doc.paragraphs[161],
        "นัยเชิงกลยุทธ์ต่อ PG: SFAS 2.48 ไม่ใช่เกณฑ์กลไกที่บังคับให้เลือก turnaround แต่เป็นหลักฐานหนึ่ง "
        "เมื่อรวมกับ core proxy ติดลบ Five Forces และช่องว่างสนาม OEM จึงสนับสนุนการฟื้นฟูพร้อมย้ายทรัพยากรไป O1/O2 "
        "โดยใช้ S2/S3/S4/S6 และแก้ W1/W3/W4 ตามลำดับ gate", size=16)

    # TOWS repaired, codes preserved, objectives separated later.
    tw = doc.tables[12]
    set_cell_text(tw.cell(0, 1), "จุดแข็ง (S)\nS1 ฐานะการเงิน\nS2 บูรณาการการผลิต\nS3 ผลิตภัณฑ์/มาตรฐาน\nS4 ฐานลูกค้าเครื่องแบบ\nS6 เครือข่ายภายใน", size=9.5, bold=True, fill=LIGHT_BLUE)
    set_cell_text(tw.cell(0, 2), "จุดอ่อน (W)\nW1 ธุรกิจหลักขาดทุน\nW2 ลูกค้า/ช่องทางกระจุกตัว\nW3 สินค้าคงเหลือ\nW4 เทคโนโลยี/ผลิตภาพ/ข้อมูล", size=9.5, bold=True, fill=LIGHT_BLUE)
    set_cell_text(tw.cell(1, 0), "โอกาส (O)\nO1 อุปสงค์ functional uniform/workwear\nO2 traceability/circularity", size=9.5, bold=True, fill=LIGHT_BLUE)
    set_cell_text(tw.cell(1, 1),
        "SO1 (S2+S3+O1) พัฒนาแนวคิด 4 แพลตฟอร์ม แต่ให้ Gate 1 เลือกเพียง 2 ที่มี margin/pilot signal สูงสุดเข้าสู่ commercial validation ก่อน\n"
        "SO2 (S3+O2) สร้างข้อมูลวัสดุ/traceability ระดับ SKU และล็อต พร้อมเอกสารผลทดสอบ\n"
        "SO3 (S4+S6+O1) ใช้ฐานลูกค้าเดิมและเครือข่ายภายในสร้าง pilot/reference customer โดยไม่ถือว่าเครือข่ายเป็น external opportunity", size=9.0)
    set_cell_text(tw.cell(1, 2),
        "WO1 (W1+O1) วิเคราะห์กำไรรายลูกค้า/SKU ครอบคลุม ≥95% ภายใน 90 วันก่อนขยาย\n"
        "WO2 (W2+O1) ตั้งทีมขาย Account-based และลดลูกค้ารายใหญ่ที่สุดเหลือ ≤30%\n"
        "WO3 (W3+W4+O2) เชื่อมข้อมูล SKU/ล็อตกับ costing และ S&OP เพื่อลด DIO ตามสูตรแบบจำลอง", size=9.0)
    set_cell_text(tw.cell(2, 0), "อุปสรรค (T)\nT1 แข่งขันด้านต้นทุน\nT2 อำนาจผู้ซื้อ\nT5 benchmark เทคโนโลยี/ความเร็วสูงขึ้น", size=9.5, bold=True, fill=LIGHT_BLUE)
    set_cell_text(tw.cell(2, 1),
        "ST1 (S2+S3+T1+T5) ถอนจาก commodity OEM และใช้ selective automation เฉพาะงานซับซ้อนที่พิสูจน์อุปสงค์/ผลตอบแทน\n"
        "ST2 (S4+S6+T2) เปลี่ยนคำสั่งซื้อครั้งเดียวเป็นสัญญาปริมาณขั้นต่ำ กลไกปรับราคา และบริการเติมสินค้า", size=9.0)
    set_cell_text(tw.cell(2, 2),
        "WT1 (W1+W3+W4+T1) ใช้ profitability hurdle, MOQ, make-to-order และ outsource งานมาตรฐาน โดย P8 ต้องผ่าน Gate 0–2\n"
        "WT2 (W2+T2) ทำ Segment and Channel Charter กับ I.C.C./บริษัทในเครือเพื่อป้องกันช่องทางทับซ้อน", size=9.0)
    set_paragraph_text(doc.paragraphs[166],
        "นัยเชิงกลยุทธ์ต่อ PG: TOWS ทั้งสิบข้อชี้ให้หยุดการแข่งขันที่ราคาเป็นตัวตัดสิน สร้างวินัยกำไรก่อนรับงาน "
        "ใช้ข้อมูลและสัญญาลดความผันผวน และลงทุนเทคโนโลยีแบบเลือกเฉพาะหลังพิสูจน์ ไม่ได้ใช้ S6 ซ้ำเป็น O ภายนอก", size=16)

    # Objectives identifiers repaired.
    obj = doc.tables[14]
    objective_texts = [
        ("OBJ1 พลิกผลจากธุรกิจหลักให้เป็นบวกอย่างยั่งยืน", "-8.02%", "คุ้มทุนปี 2571 และ ≥3% ปี 2572", "≥7%", "W1, WO1, WT1"),
        ("OBJ2 ยกระดับอัตรากำไรขั้นต้นด้วยส่วนผสมผลิตภัณฑ์ใหม่", "21.40%", "≥24.0% ปี 2571", "≥27.5%", "SO1, ST1, ST2"),
        ("OBJ3 ลดค่าใช้จ่ายขายและบริหารเชิงโครงสร้าง", "29.42%", "≤23.0% ปี 2571", "≤20.0%", "W1, WT1"),
        ("OBJ4 ลดการพึ่งพาลูกค้ารายใหญ่ที่สุด", "40.12%", "≤35% ปี 2572", "≤30%", "W2, WO2, WT2"),
        ("OBJ5 ลดเงินทุนที่ตรึงในสินค้าคงเหลือ", "DIO แบบจำลอง 267.75 วัน", "≤210 วันปี 2571 และ ≤170 วันปี 2572", "ประมาณ 110 วัน; CCC ≤125 วัน", "W3, W4, WO3"),
        ("OBJ6 เพิ่มสัดส่วนรายได้ functional uniform/workwear", "กำหนดฐานใน 90 วัน", "≥30% ปี 2572", "≥38%", "O1, SO1, WO1"),
        ("OBJ7 ยกระดับผลตอบแทนต่อส่วนของผู้ถือหุ้น", "-0.40%", "≥4% ปี 2572", "≥6%", "W5, S1"),
        ("OBJ8 สร้างข้อมูลตรวจสอบย้อนกลับสำหรับสินค้าเชิงกลยุทธ์", "ต้องกำหนดฐาน", "minimum viable compliance ก่อน 14 ธ.ค. 2570; ครบ 100% ปี 2572", "รักษา 100%", "O2, SO2"),
        ("OBJ9 เติบโตของยอดขายอย่างมีคุณภาพ", "605.20 ล้านบาท", "≥700 ล้านบาท ปี 2572", "ประมาณ 810 ล้านบาท", "CAGR โดยนัย 4.98%; ไม่ใช่ company guidance"),
    ]
    for ri, vals in enumerate(objective_texts, start=1):
        for ci, val in enumerate(vals):
            set_cell_text(obj.cell(ri, ci), val, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT if ci in {0,4} else WD_ALIGN_PARAGRAPH.CENTER)
    set_paragraph_text(doc.paragraphs[172],
        "นัยเชิงกลยุทธ์ต่อ PG: OBJ9 ใช้ยอดขายปี 2574 ประมาณ 810 ล้านบาท ซึ่งให้ CAGR โดยนัยจากปี 2568 ร้อยละ 4.98 "
        "อัตรานี้อยู่ในช่วงเดียวกับกรอบการเติบโตร้อยละ 4–5 ต่อปีที่บริษัทเคยประกาศในคนละ horizon "
        "จึงไม่ตีความว่า 810 ล้านบาทเป็นเป้าหมายที่บริษัทประกาศ", size=16)

    # Structure 4.4–4.7 and rescored alternatives.
    set_paragraph_text(doc.paragraphs[173], "4.4 ทางเลือกเชิงกลยุทธ์ (Strategic Alternatives)", style="Heading 2", size=16, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[174], "4.4.1 ทางเลือก A/B/C และเกณฑ์พิจารณา", style="Heading 3", size=15, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[176],
        "ทางเลือก A: Volume-led Export OEM เร่งปริมาณงานรับจ้างผลิตเพื่อส่งออกและลงทุน automation เพื่อไล่ตามขนาดตลาด "
        "OIE รายงานการส่งออกปี 2568 ประมาณ 2,018.77 ล้านดอลลาร์สหรัฐ +5.61% จึงให้น้ำหนักความน่าดึงดูดตลาดสูงขึ้น "
        "แต่ PG ยังเสียเปรียบขนาด ต้นทุนแรงงาน และฐานผลิตหลายประเทศ และยอดส่งออกของ PG ลด 34.36% ทางเลือกนี้จึงเสี่ยงต่อการกลับไปแข่ง commodity OEM", size=16)
    set_paragraph_text(doc.paragraphs[177],
        "ทางเลือก B: Own-brand D2C Omnichannel สร้างแบรนด์ ช่องทางตรง และข้อมูลลูกค้า Sabina/MC Group ใช้เป็น comparables ด้าน brand/channel economics เท่านั้น "
        "ไม่ใช่ฐานสมมติว่า PG จะได้อัตรากำไรขั้นต้นเท่ากัน ทางเลือกนี้ช่วยลดการพึ่งผู้จัดจำหน่าย แต่ต้องสร้างความสามารถแบรนด์ การตลาดตรง และการบริหารสินค้าคงเหลือ "
        "ซึ่งจากข้อมูลสาธารณะที่ตรวจสอบยังไม่พบหลักฐานว่ามีขนาดพร้อมรองรับการเปลี่ยนโมเดลทั้งบริษัท", size=16)
    set_paragraph_text(doc.paragraphs[178],
        "ทางเลือก C: Focused Functional Uniform & Workwear Solutions ฟื้นกำไรธุรกิจหลักและมุ่งลูกค้าองค์กรที่ซื้อด้วยคุณสมบัติ/มาตรฐาน "
        "ใช้ฐาน S2/S3/S4/S6 และสัญญาเติมสินค้า แต่มีวงจรขาย B2B ยาว ความเสี่ยง channel conflict และภาระ compliance ก่อน 14 ธันวาคม 2570 "
        "จึงลดคะแนนความเสี่ยงในการนำไปปฏิบัติจาก 3 เป็น 2 และบังคับ Gate 0–2 ก่อน scale/capex", size=16)
    set_paragraph_text(doc.paragraphs[179],
        "เกณฑ์ทั้งเจ็ดและน้ำหนักกำหนดก่อนให้คะแนน โดยให้น้ำหนักสูงสุดแก่ความสอดคล้องกับทรัพยากรและศักยภาพอัตรากำไร "
        "(ข้อละ 0.20) เพราะปัญหาหลักคือ core proxy ติดลบ ส่วนเงินลงทุน ความเสี่ยงการปฏิบัติ และความเร็วสู่ผลลัพธ์มีน้ำหนักข้อละ 0.10 "
        "ฐานะการเงินช่วยลดข้อจำกัดด้านเงินทุน แต่ไม่ทดแทนความสามารถที่ต้องสร้างใหม่หรือหลักฐาน product-market fit", size=16)
    insert_paragraph_before(doc, doc.paragraphs[180], "4.5 เมทริกซ์การตัดสินใจและการทดสอบความไว", style="Heading 2", size=16, bold=True, color=BLUE)
    dm = doc.tables[15]
    decision_rows = [
        ["เกณฑ์การตัดสินใจ", "น้ำหนัก", "A คะแนน", "A ถ่วง", "B คะแนน", "B ถ่วง", "C คะแนน", "C ถ่วง"],
        ["1. สอดคล้องทรัพยากร/ความสามารถ", "0.20", "3", "0.60", "2", "0.40", "5", "1.00"],
        ["2. ความน่าดึงดูดของตลาด", "0.15", "4", "0.60", "4", "0.60", "4", "0.60"],
        ["3. ศักยภาพอัตรากำไร", "0.20", "1", "0.20", "4", "0.80", "4", "0.80"],
        ["4. เงินลงทุน (สูง = ลงทุนน้อย)", "0.10", "2", "0.20", "2", "0.20", "4", "0.40"],
        ["5. การแข่งขัน/ความยากเลียนแบบ", "0.15", "1", "0.15", "2", "0.30", "4", "0.60"],
        ["6. ความเสี่ยงการนำไปปฏิบัติ (สูง = เสี่ยงต่ำ)", "0.10", "2", "0.20", "2", "0.20", "2", "0.20"],
        ["7. ความเร็วสู่ผลลัพธ์", "0.10", "2", "0.20", "2", "0.20", "3", "0.30"],
        ["รวม (กรณีฐาน)", "1.00", "", "2.15", "", "2.70", "", "3.90"],
        ["การทดสอบความไว", "", "", "", "", "", "", ""],
        ["กรณี 1: ตลาด 0.30; ความสอดคล้อง 0.05", "1.00", "", "2.30", "", "3.00", "", "3.75"],
        ["กรณี 2: ลดคะแนน margin ของ C จาก 4 เป็น 2", "1.00", "", "2.15", "", "2.70", "", "3.50"],
        ["กรณี 3: B ได้ตลาดและ margin = 5", "1.00", "", "2.15", "", "3.05", "", "3.90"],
    ]
    rebuild_table(dm, decision_rows, section_rows={9}, widths=[3000, 700, 650, 700, 650, 700, 650, 700],
                  font_size=9.5, numeric_cols={1,2,3,4,5,6,7}, emphasis_rows={8,10,11,12})
    set_paragraph_text(doc.paragraphs[183],
        "นัยเชิงกลยุทธ์ต่อ PG: C ยังชนะหลังลดคะแนนความเสี่ยง และ A ได้คะแนนตลาดสูงขึ้นตามข้อมูล OIE "
        "จึงไม่ได้ออกแบบคะแนนเพื่อให้ C ชนะ ความได้เปรียบของ C มาจาก resource fit, margin potential, investment discipline "
        "และการแข่งขันที่หลีกเลี่ยง scale race ส่วน sensitivity test แสดงว่าผลไม่ขึ้นกับสมมติฐานเดียว", size=16)
    set_paragraph_text(doc.paragraphs[184], "4.6 กลยุทธ์ที่แนะนำ (Recommended Strategy)", style="Heading 2", size=16, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[186],
        "นัยเชิงกลยุทธ์ต่อ PG: กลยุทธ์ที่แนะนำคือ Focused Turnaround + Concentrated Growth ภายใต้ Focused Differentiation "
        "มุ่ง functional uniform/workwear/medical และ selected high-value ODM โดยพัฒนา 4 concepts แต่ให้เพียง 2 แพลตฟอร์มที่ผ่าน Gate 1 "
        "เข้าสู่ commercial validation ก่อน เพื่อลดการกระจายทรัพยากร", size=16)
    set_paragraph_text(doc.paragraphs[185],
        "ประเด็นเชิงกลยุทธ์: PG มีฐานะการเงินแข็งแรง มาตรฐานภายนอก การผลิตแบบบูรณาการ และเครือข่ายผู้ถือหุ้นเป็นฐาน "
        "แต่ core proxy ติดลบต่อเนื่อง ขณะที่ลูกค้า/ช่องทางกระจุกตัวและการปรับค่าใช้จ่ายตามยอดขายทำได้ช้า จากข้อมูลสาธารณะที่ตรวจสอบ "
        "ไม่พบหลักฐานเพียงพอว่ามีกลไก profitability gate รายลูกค้า/SKU จึงต้องเลือกสนามที่คุณสมบัติและบริการสร้างมูลค่าได้มากกว่าต้นทุน "
        "พร้อมจัดโครงสร้างผู้รับผิดชอบและทุนใหม่ ข้อสรุปนี้เป็นการอนุมานจาก [1]–[4] และตารางที่ 7–16", size=16)
    p187 = doc.paragraphs[187]
    insert_paragraph_before(doc, p187, "4.7 กลยุทธ์ระดับองค์กร ธุรกิจ และหน้าที่", style="Heading 2", size=16, bold=True, color=BLUE)
    set_paragraph_text(p187, "4.7.1 กลยุทธ์ระดับองค์กร (Corporate Strategy)", style="Heading 3", size=15, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[189], "4.7.2 กลยุทธ์ระดับธุรกิจ (Business Strategy)", style="Heading 3", size=15, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[190],
        "กลยุทธ์ระดับธุรกิจคือการสร้างความแตกต่างแบบมุ่งเน้น (Focused Differentiation) ไม่เลือก cost leadership เพราะ PG มีขนาด "
        "2.712 ล้านชิ้น เทียบ benchmark Hi-Tech มากกว่า 40 ล้านชิ้นและไม่มีฐานผลิตหลายประเทศ [1][14] ไม่เลือก broad differentiation "
        "เพราะต้องสร้างแบรนด์/ช่องทางตรงในวงกว้างทั้งที่ความพร้อมดังกล่าวยังไม่ผ่านการพิสูจน์ เลือก focus เพื่อลดการเปรียบเทียบราคาตรงและเพิ่ม switching cost "
        "ด้วยข้อกำหนด/ผลทดสอบรายแพลตฟอร์ม traceability ระดับ SKU/ล็อต SLA การส่งมอบ/เติมสินค้า และข้อมูลต้นทุนตลอดอายุการใช้งาน "
        "ทุกฐานความแตกต่างต้องพิสูจน์ใน Gate 1 ไม่ถือเป็นข้อเท็จจริงล่วงหน้า", size=16)
    set_paragraph_text(doc.paragraphs[191], "4.7.3 กลยุทธ์ระดับหน้าที่ (Functional Strategy)", style="Heading 3", size=15, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[195],
        "การพัฒนาผลิตภัณฑ์และวิจัย พัฒนาแนวคิด 4 แพลตฟอร์ม แต่จัดลำดับ 2 แพลตฟอร์มที่มี profitability/pilot signal สูงสุด "
        "สำหรับการพิสูจน์เชิงพาณิชย์รอบแรก ทุกแพลตฟอร์มต้องมีข้อกำหนด ผลทดสอบ และข้อมูลอายุการใช้งานก่อนขยาย (SO1, SO2)", size=16)
    set_paragraph_text(doc.paragraphs[199],
        "ห่วงโซ่อุปทานและความยั่งยืน เริ่ม P7 ตั้งแต่ปลาย 2569 จัดทำ minimum viable compliance ครอบคลุมสินค้ากลยุทธ์เสี่ยงสูงก่อน 14 ธันวาคม 2570 "
        "และติดตามกฎเฉพาะผลิตภัณฑ์ของ ESPR/DPP แทนการสมมติกำหนดวันบังคับสำหรับเสื้อผ้าทั้งหมด (SO2)", size=16)
    set_paragraph_text(doc.paragraphs[201],
        "นัยเชิงกลยุทธ์ต่อ PG: กลยุทธ์สามระดับผูกด้วย Gate 0–4 หาก P1 ไม่ยืนยัน margin หรือ pilot ไม่ยืนยัน demand ให้คง turnaround "
        "และชะลอ growth/capex สิ่งที่ PG จะไม่ทำคือ scale race ใน commodity OEM, capacity expansion ที่ไม่มี committed demand/business case, "
        "full consumer-brand transformation ระหว่างแผน และงาน tender/order ที่ไม่ผ่าน profitability hurdle", size=16)

    # Action plan with transition period, explicit gates and corrected P8 basis.
    set_paragraph_text(doc.paragraphs[204],
        "แผนหลักครอบคลุม 2570–2574 และมีช่วง Transition/Readiness ตั้งแต่ไตรมาส 4/2569 สำหรับ P7 และการเตรียมข้อมูล Gate 0 "
        "Gate 0 = Data and profitability validation; Gate 1 = Pilot/product-market validation; Gate 2 = Commercial scaling; "
        "Gate 3 = Selective automation/capex; Gate 4 = Selective export expansion P1 เป็นเงื่อนไขแรกและ P8 ห้ามอนุมัติก่อนมีผล P1/pilot/committed demand/business case", size=16)
    set_paragraph_text(doc.paragraphs[205],
        "ลำดับปี: 2570 วางฐาน/พิสูจน์และเลือก 2 แพลตฟอร์มแรก; 2571 เปิดตลาดและผ่าน Gate 2; 2572 ขยายผลและพิจารณา Gate 3; "
        "2573 เพิ่มประสิทธิภาพ; 2574 ขยายส่งออกแบบเลือกเฉพาะผ่าน Gate 4 ขณะที่ P7 เริ่มปลาย 2569 และต้องมี minimum viable compliance ก่อน 14 ธันวาคม 2570", size=16)

    ap = doc.tables[16]
    # Targeted cell edits keep the original action-plan geometry.
    action_updates = {
        (2,1): "P1 วินิจฉัยความสามารถทำกำไรรายลูกค้าและราย SKU (Gate 0)",
        (2,2): "จัดทำ core P&L แยกจากรายได้อื่น วิเคราะห์ margin รายลูกค้า/SKU/ช่องทางย้อนหลัง 24 เดือน ครอบคลุม ≥95%; เป็น prerequisite ของ P2–P10",
        (3,2): "ดำเนินหลัง Gate 0: ยุติ/เจรจาใหม่งานที่ไม่ผ่าน profitability hurdle กำหนด MOQ และทบทวนค่าใช้จ่าย/การจัดสรรทุน",
        (4,2): "จัดตั้งหน่วยธุรกิจ P&L และทีม Account-based; อนุมัติ scale เมื่อผ่าน Gate 0 และ pipeline ผ่าน Gate 1",
        (5,1): "P4 พัฒนา 4 concepts และคัด 2 แพลตฟอร์มแรก (Gate 1)",
        (5,2): "พัฒนา 4 concepts (medical, service/food, industrial, circular corporate) พร้อม spec/test; เลือก 2 ที่ margin/pilot signal สูงสุดเข้าสู่ commercial validation",
        (5,7): "4 concepts มีเอกสารขั้นต่ำ; 2 แพลตฟอร์มผ่าน Gate 1; conversion ≥25%",
        (6,2): "pilot กับองค์กรในเครือ โรงพยาบาล และโรงแรม; เก็บ usage/margin/reorder evidence เพื่อ Gate 1",
        (8,2): "ติดตั้ง CRM, digital costing และ S&OP หลังข้อมูล Gate 0 พร้อมใช้งาน; ระบบต้องบังคับ profitability hurdle ก่อนเสนอราคา",
        (9,1): "P7 traceability/material data readiness",
        (9,2): "เริ่มปลาย 2569 ทำ supplier/material/SKU-lot data และ forced-labour due diligence; minimum viable compliance ก่อน 14 ธ.ค. 2570; ติดตามกฎเฉพาะผลิตภัณฑ์ ESPR/DPP",
        (9,4): "ปลาย 2569–2572",
        (9,7): "MVP สินค้าเสี่ยงสูง 100% ก่อน 14 ธ.ค. 2570; SKU เชิงกลยุทธ์ 100% ปี 2572",
        (11,1): "P8 selective automation/capex (Gate 3)",
        (11,2): "อนุมัติทีละสายหลัง P1, Gate 1–2, committed demand และ business case; pilot หนึ่งสาย วัด labor time/OEE/OTD ก่อนขยาย",
        (11,4): "ปี 2572–2574",
        (13,2): "พอร์ทัลเติมสินค้าหลัง Gate 2 เฉพาะลูกค้าที่มีสัญญา/SLA และข้อมูล S&OP พร้อม",
        (14,1): "P10 selected high-value ODM export (Gate 4)",
        (14,2): "ขยายเฉพาะ ODM ที่มีข้อกำหนดสูง margin ผ่าน hurdle และ traceability พร้อม; ไม่กลับไปแข่ง volume-led commodity OEM",
        (16,2): "สรรหา/พัฒนา sales engineering, textile technology, IE/data และปกป้องตำแหน่งวิกฤตระหว่าง turnaround; เป็น enabling prerequisite ทุก Gate",
    }
    for (ri, ci), value in action_updates.items():
        set_cell_text(ap.cell(ri, ci), value, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_paragraph_text(doc.paragraphs[210], "5.2 งบประมาณและฐานการประมาณการ", style="Heading 2", size=16, bold=True, color=BLUE)
    bud = doc.tables[17]
    set_cell_text(bud.cell(4,2), "ค่าพัฒนา 4 concepts และการทดสอบ โดยให้เพียง 2 แพลตฟอร์มผ่าน Gate 1 รอบแรก; ไม่ลงทุนโครงสร้างพื้นฐานใหม่ก่อน validation", size=9.5)
    set_cell_text(bud.cell(7,2), "เริ่มปลาย 2569 ต่อเนื่องจากระบบคู่ค้า 222 ราย ครอบคลุม supplier/material/SKU-lot data และ forced-labour readiness ก่อน 14 ธ.ค. 2570", size=9.5)
    set_cell_text(bud.cell(8,2),
        "คงกรอบ 45–75 ล้านบาทตลอด 3 ปี เท่ากับเฉลี่ย 15–25 ล้านบาท/ปี หรือประมาณ 0.94–1.57 เท่าของเงินลงทุนเครื่องจักรปี 2568 ที่ 15.96 ล้านบาทต่อปี "
        "ไม่ใช่ 2–3 เท่าต่อปี อนุมัติเป็นขั้นเฉพาะสายที่ผ่าน Gate 0–2", size=9.5)
    set_paragraph_text(doc.paragraphs[212],
        "นัยเชิงกลยุทธ์ต่อ PG: กรอบงบ 102.5–174.0 ล้านบาทเป็นสมมติฐานเชิงวางแผน ไม่ใช่งบที่บริษัทประกาศ "
        "ฉากทัศน์ให้เงินทุนมาจากกระแสเงินสดดำเนินงาน รายได้จากเงินลงทุน และการครบกำหนด/ปรับสัดส่วนสินทรัพย์ทางการเงินแบบเลือกเฉพาะ "
        "โดยไม่ต้องเพิ่มทุนหรือก่อหนี้ใหม่ ทั้งนี้ไม่สมมติว่าต้องขายสินทรัพย์ทางการเงินหลัก และทุกโครงการต้องมีใบเสนอราคา/business case ก่อนอนุมัติ", size=16)
    set_paragraph_text(doc.paragraphs[213], "5.3 Balanced Scorecard และการควบคุมเชิงกลยุทธ์", style="Heading 2", size=16, bold=True, color=BLUE)
    bsc = doc.tables[18]
    bsc_repl = {
        2: ("OBJ1", "ผลจากธุรกิจหลัก"),
        3: ("OBJ2 และ OBJ3", "อัตรากำไรขั้นต้น/SG&A ต่อรายได้"),
        4: ("OBJ7 และ OBJ9", "ROE ใช้กำไรสุทธิ/ส่วนของผู้ถือหุ้นเฉลี่ย และยอดขาย"),
        6: ("OBJ4", "สัดส่วนลูกค้ารายใหญ่ที่สุด"),
        7: ("OBJ6", "สัดส่วนรายได้ Functional Workwear"),
        10: ("OBJ5", "DIO = สินค้าคงเหลือเฉลี่ย/COGS×365; CCC = AR days + DIO − AP days"),
        13: ("OBJ8", "สัดส่วน SKU เชิงกลยุทธ์ที่มีข้อมูล traceability ครบ"),
    }
    for ri, (objective, kpi_text) in bsc_repl.items():
        set_cell_text(bsc.cell(ri,1), objective, size=8.5)
        set_cell_text(bsc.cell(ri,2), kpi_text, size=8.5)
    set_cell_text(bsc.cell(10,3), "DIO แบบจำลอง 267.75 วัน; CCC แบบจำลองประมาณ 287.59 วัน", size=8.5)
    set_cell_text(bsc.cell(10,4), "DIO ประมาณ 110 วัน; CCC ประมาณ 120 วัน", size=8.5)
    set_cell_text(bsc.cell(13,4), "MVP ก่อน 14 ธ.ค. 2570; ครบ 100% ปี 2572 และรักษาระดับ", size=8.5)
    set_paragraph_text(doc.paragraphs[216],
        "นัยเชิงกลยุทธ์ต่อ PG: รายเดือนติดตาม core P&L, margin รายลูกค้า/SKU, DIO/CCC ตามสูตรแบบจำลอง, OTD และ pipeline "
        "รายไตรมาสทบทวน SFAS, risk register และ Gate decisions รายปีทบทวน strategy/objectives/competitor benchmark "
        "ROE ในฉากทัศน์ใช้ส่วนของผู้ถือหุ้นเฉลี่ยที่สร้างจากงบแสดงฐานะการเงิน ไม่ใช่ยอดปลายงวดโดยไม่ระบุฐาน", size=16)
    set_paragraph_text(doc.paragraphs[217], "5.4 Risk Register / Early Warning / Contingency", style="Heading 2", size=16, bold=True, color=BLUE)

    # Financial model: correct formulas, remove unsupported base-case ROE, add full scenario assumptions and SFP.
    base_rev = [605.20, 592.00, 586.00, 580.00, 574.00, 568.00]
    base_gm = [21.40, 21.00, 20.80, 20.60, 20.40, 20.20]
    base_gp = [129.52] + [base_rev[i] * base_gm[i] / 100 for i in range(1,6)]
    base_sga = [178.08, 152.00, 151.00, 150.00, 149.00, 148.00]
    base_core = [base_gp[i]-base_sga[i] for i in range(6)]
    base_other = [62.14, 60,60,60,60,60]
    base_ni = [-5.56] + [(base_core[i]+base_other[i])*0.8 for i in range(1,6)]
    base_inv = [337.11,340,342,344,346,348]
    base_cogs = [base_rev[i]-base_gp[i] for i in range(6)]
    base_dio = [267.75]
    prev = 337.105281
    for i in range(1,6):
        end = base_inv[i]
        base_dio.append(((prev+end)/2)/base_cogs[i]*365)
        prev = end

    p222 = doc.paragraphs[222]
    set_paragraph_text(p222, "6.1 กรณีฐาน (Base Case)", style="Heading 2", size=16, bold=True, color=BLUE)
    insert_paragraph_after(doc, p222, "ตารางที่ 21 ประมาณการผลการดำเนินงาน กรณีฐาน ปี 2570–2574 (หน่วย: ล้านบาท เว้นแต่ระบุ)", size=14, bold=True)
    base_rows = [
        ["รายการ", "2568 (จริง)", "2570", "2571", "2572", "2573", "2574"],
        ["ยอดขาย"] + [money(x) for x in base_rev],
        ["อัตรากำไรขั้นต้น (%)"] + [pct(x) for x in base_gm],
        ["กำไรขั้นต้น"] + [money(x) for x in base_gp],
        ["ค่าใช้จ่ายขายและบริหาร"] + [money(x) for x in base_sga],
        ["SG&A/ยอดขาย (%)"] + [pct(base_sga[i]/base_rev[i]*100) for i in range(6)],
        ["ผลจากธุรกิจหลักโดยประมาณ"] + [f"({abs(x):.2f})" if x < 0 else money(x) for x in base_core],
        ["อัตราผลจากธุรกิจหลัก (%)"] + [f"({abs(base_core[i]/base_rev[i]*100):.2f})" if base_core[i] < 0 else pct(base_core[i]/base_rev[i]*100) for i in range(6)],
        ["รายได้อื่น"] + [money(x) for x in base_other],
        ["กำไรสุทธิโดยประมาณ (ภาษี 20%)*"] + [f"({abs(x):.2f})" if x < 0 else money(x) for x in base_ni],
        ["สินค้าคงเหลือปลายงวด"] + [money(x) for x in base_inv],
        ["DIO: สินค้าคงเหลือเฉลี่ย/COGS×365 (วัน)"] + [pct(x) for x in base_dio],
    ]
    rebuild_table(doc.tables[20], base_rows, widths=[2600,1070,1070,1070,1070,1070,1070], font_size=9.5,
                  numeric_cols={1,2,3,4,5,6}, emphasis_rows={6,7})
    set_paragraph_text(doc.paragraphs[224],
        "นัยเชิงกลยุทธ์ต่อ PG: กรณีฐานยังแสดงกำไรสุทธิจากรายได้อื่น ขณะที่ core proxy ติดลบเพิ่มเป็นประมาณ 33.26 ล้านบาทในปี 2574 "
        "ไม่แสดง ROE เพราะไม่ได้สร้างงบแสดงฐานะการเงินของกรณีฐานครบถ้วน DIO คำนวณจาก average inventory/COGS และเพิ่มจากประมาณ 267.75 เป็น 279.43 วัน "
        "จึงสะท้อนการตรึงทุนต่อเนื่องโดยไม่ใช้วิธี scaling ยอดคงเหลือด้วยยอดขาย", size=16)

    # Strategic scenario and reconciled equity/ROE.
    rev = [600.00,645.00,700.00,755.00,810.00]
    growth = [-0.86,7.50,8.53,7.86,7.28]
    fwork = [130,175,225,275,310]
    gm = [22.5,24.0,25.5,26.5,27.5]
    gp = [rev[i]*gm[i]/100 for i in range(5)]
    cogs = [rev[i]-gp[i] for i in range(5)]
    sga = [150,148.35,154,158.55,162]
    core = [gp[i]-sga[i] for i in range(5)]
    other = [58,56,54,54,54]
    ni = [(core[i]+other[i])*0.8 for i in range(5)]
    base_forecast_ni = base_ni[1:]
    inventory = [300,260,220,190,165]
    inventory_dio = []
    prev_inv = 337.105281
    for i in range(5):
        inventory_dio.append(((prev_inv+inventory[i])/2)/cogs[i]*365)
        prev_inv = inventory[i]

    opening_equity_2570 = 1365.541175 + 30.0 - 48.0 + 25.14
    dividends = [0,20,30,40,50]
    equity = []
    e = opening_equity_2570
    for n,dv in zip(ni,dividends):
        e = e + n - dv
        equity.append(e)
    roe = []
    for i,n in enumerate(ni):
        opening = opening_equity_2570 if i == 0 else equity[i-1]
        roe.append(n/((opening+equity[i])/2)*100)

    p225 = doc.paragraphs[225]
    set_paragraph_text(p225, "6.2 กรณีกลยุทธ์ (Strategic Case)", style="Heading 2", size=16, bold=True, color=BLUE)
    insert_paragraph_after(doc, p225, "ตารางที่ 22 ประมาณการผลการดำเนินงาน กรณีกลยุทธ์ ปี 2570–2574 (หน่วย: ล้านบาท เว้นแต่ระบุ)", size=14, bold=True)
    strategic_rows = [
        ["รายการ", "2568 (จริง)", "2570", "2571", "2572", "2573", "2574"],
        ["ยอดขาย", "605.20"] + [money(x) for x in rev],
        ["อัตราเติบโตยอดขาย (%)", "(21.80)"] + [pct(x) if x>=0 else f"({abs(x):.2f})" for x in growth],
        ["รายได้ functional uniform/workwear", "ต้องกำหนดฐาน"] + [money(x) for x in fwork],
        ["สัดส่วนยอดขายรวม (%)", "ต้องกำหนดฐาน"] + [pct(fwork[i]/rev[i]*100) for i in range(5)],
        ["อัตรากำไรขั้นต้น (%)", "21.40"] + [pct(x) for x in gm],
        ["กำไรขั้นต้น", "129.52"] + [money(x) for x in gp],
        ["ค่าใช้จ่ายขายและบริหาร", "178.08"] + [money(x) for x in sga],
        ["SG&A/ยอดขาย (%)", "29.42"] + [pct(sga[i]/rev[i]*100) for i in range(5)],
        ["ผลจากธุรกิจหลักโดยประมาณ", "(48.56)"] + [f"({abs(x):.2f})" if x<0 else money(x) for x in core],
        ["อัตราผลจากธุรกิจหลัก (%)", "(8.02)"] + [f"({abs(core[i]/rev[i]*100):.2f})" if core[i]<0 else pct(core[i]/rev[i]*100) for i in range(5)],
        ["รายได้อื่น", "62.14"] + [money(x) for x in other],
        ["กำไรสุทธิโดยประมาณ (ภาษี 20%)*", "(5.56)"] + [money(x) for x in ni],
        ["สินค้าคงเหลือปลายงวด", "337.11"] + [money(x) for x in inventory],
        ["DIO: สินค้าคงเหลือเฉลี่ย/COGS×365 (วัน)", "267.75"] + [pct(x) for x in inventory_dio],
        ["ROE: กำไรสุทธิ/ส่วนของผู้ถือหุ้นเฉลี่ย (%)", "(0.40)"] + [pct(x) for x in roe],
        ["ส่วนต่างกำไรสุทธิเทียบกรณีฐาน", "–"] + [money(ni[i]-base_forecast_ni[i]) for i in range(5)],
    ]
    rebuild_table(doc.tables[21], strategic_rows, widths=[2600,1070,1070,1070,1070,1070,1070], font_size=9.2,
                  numeric_cols={1,2,3,4,5,6}, emphasis_rows={9,10,15})

    p229 = doc.paragraphs[229]
    set_paragraph_text(p229, "6.3 สมมติฐานของประมาณการ", style="Heading 2", size=16, bold=True, color=BLUE)
    insert_paragraph_after(doc, p229, "ตารางที่ 23 ตารางสมมติฐานฉากทัศน์เป้าหมายเชิงบริหาร (Management Target Scenario – not company guidance)", size=14, bold=True)

    ar_days = [58,56,54,52,50]
    ap_days = [38,39,40,40,40]
    ar = [rev[i]*ar_days[i]/365 for i in range(5)]
    ap = [cogs[i]*ap_days[i]/365 for i in range(5)]
    capex = [16,16,36,41,31]
    depreciation = [12,13,15,18,20]
    ppe = []
    ppe_balance = 165.17946
    for c,d in zip(capex,depreciation):
        ppe_balance = ppe_balance + c - d
        ppe.append(ppe_balance)
    cash = [25,30,35,40,45]
    lease_total = [16,14,12,10,8]
    current_lease = [8,7,6,5,4]
    noncurrent_lease = [8,7,6,5,4]
    other_current_liab = [6.3]*5
    employee_benefit = [50,51,52,53,54]
    deferred_tax = [22.7,20.7,18.7,16.7,14.7]
    total_liab = [ap[i]+lease_total[i]+other_current_liab[i]+employee_benefit[i]+deferred_tax[i] for i in range(5)]
    total_assets = [total_liab[i]+equity[i] for i in range(5)]
    other_current_assets = [7.3]*5
    other_noncurrent_assets = [0.39 + noncurrent_lease[i] + 0.47 + 13.81 for i in range(5)]
    financial_assets = []
    current_financial = []
    noncurrent_financial = []
    for i in range(5):
        fa = total_assets[i] - (cash[i]+ar[i]+inventory[i]+other_current_assets[i]+ppe[i]+other_noncurrent_assets[i])
        financial_assets.append(fa)
        current_financial.append(fa*0.275)
        noncurrent_financial.append(fa*0.725)

    assump_rows = [
        ["สมมติฐาน", "ค่าที่ใช้ 2570–2574", "ฐาน/สูตรและข้อจำกัด", "เชื่อมโยง"],
        ["รายได้", "600 / 645 / 700 / 755 / 810", "สมมติฐานเชิงวางแผนจาก mix; CAGR 2568–2574 = 4.98% อยู่ในช่วงเดียวกับกรอบเดิมของบริษัท แต่ไม่ใช่ company guidance", "P1–P5, P9–P10"],
        ["COGS", "465.00 / 490.20 / 521.50 / 554.93 / 587.25", "Revenue × (1 − Gross Margin); Gross Profit = Revenue × Gross Margin", "P2, P4, P8"],
        ["AR days", "58 / 56 / 54 / 52 / 50", "Receivables = Revenue × AR days/365; ใช้ trade and other receivables เป็น proxy เพราะไม่พบการแยกอื่นเพียงพอ", "P3, P6"],
        ["Inventory / DIO", "Ending inventory 300 / 260 / 220 / 190 / 165; DIO 250.05 / 208.49 / 167.98 / 134.84 / 110.32", "DIO = Average Inventory/COGS×365 โดย Average = (Beginning+Ending)/2; ไม่ scaling ending inventory ด้วยยอดขาย", "P2, P6, P9"],
        ["AP days", "38 / 39 / 40 / 40 / 40", "Trade payables = COGS × AP days/365; เป็น management proxy", "Supply chain"],
        ["Capex", "16 / 16 / 36 / 41 / 31", "maintenance ใกล้ฐานเครื่องจักรปี 2568; selective P8 ส่วนเพิ่มเฉพาะ 2572–2574 หลัง Gate 3", "P8"],
        ["Depreciation", "12 / 13 / 15 / 18 / 20", "สมมติฐานเชิงวางแผนสัมพันธ์กับ opening PPE และ capex; ต้องแทนด้วย asset register จริงก่อนอนุมัติ", "P8"],
        ["Financial assets", "ยอดปลายงวดเป็น residual liquidity; current/non-current = 27.5%/72.5%", "อาจใช้กระแสเงินสด รายได้ลงทุน และ selective maturity/reallocation โดยไม่บังคับขายหลักทรัพย์; residual เพิ่มได้เมื่อกำไร/ลด inventory สูงกว่าการใช้เงิน", "Finance"],
        ["Debt / lease", "ไม่มี bank debt ใหม่; lease total 16 / 14 / 12 / 10 / 8", "planning assumption; ไม่มี new equity/new debt ภายใต้ scenario", "Gate 3"],
        ["Dividends", "0 / 20 / 30 / 40 / 50", "สมมติให้พักปันผลปีแรกและจ่ายแบบมีเงื่อนไขตาม liquidity/gate; ไม่ใช่นโยบายบริษัท", "Board decision"],
        ["Retained earnings", "กำไรสะสมไม่จัดสรรเพิ่มด้วย Net Profit − Dividend", "bridge 2569 ใช้ FY2569 net profit assumption 30, dividend paid 48, H1 OCI +25.14 และไม่สมมติ OCI เพิ่มใน H2", "SFP / ROE"],
        ["Cash/liquidity", "25 / 30 / 35 / 40 / 45", "minimum management cash buffer; financial assets เป็น residual หลังรักษา buffer", "Finance"],
        ["Other income", "58 / 56 / 54 / 54 / 54", "conservative yield/mix assumption; ไม่ตีความว่าต้องขาย financial assets และไม่รวม valuation gain", "Income statement"],
        ["Tax / net profit", "20% ของ core proxy + other income", "ประมาณการเชิงบริหาร ไม่ใช่ audited tax forecast; ไม่รวม valuation gain/loss", "Strategic case"],
    ]
    rebuild_table(doc.tables[22], assump_rows, widths=[1700,1900,4200,1220], font_size=9.0)
    set_paragraph_text(doc.paragraphs[230],
        "นัยเชิงกลยุทธ์ต่อ PG: สมมติฐาน inventory ถูกสร้างจาก COGS และค่าเฉลี่ย ไม่แสดง claim ว่าจะปล่อยเงิน 181 ล้านบาทจากการ scaling "
        "ยอดคงเหลือแบบเดิม ส่วนต่างสินค้าคงเหลือปลายงวด 337.11 เป็น 165.00 ล้านบาทเป็นเพียงผลของ management target และต้องยืนยันด้วยข้อมูล SKU/S&OP "
        "การจัดหาเงินทุนไม่ขัดแย้งกัน: scenario อาจใช้ operating cash, investment income และ selective maturity/reallocation โดยไม่บังคับขายสินทรัพย์หลัก", size=16)

    # Insert projected statement of financial position before limitations.
    p231 = doc.paragraphs[231]
    h64 = insert_paragraph_before(doc, p231, "6.4 งบแสดงฐานะการเงินประมาณการ", style="Heading 2", size=16, bold=True, color=BLUE)
    cap64 = insert_paragraph_after(doc, h64,
        "ตารางที่ 24 งบแสดงฐานะการเงินประมาณการ: Management Target Scenario – not company guidance (หน่วย: ล้านบาท)",
        size=14, bold=True)
    sfp = add_table_after(doc, cap64, rows=1, cols=7, widths=[2600,1070,1070,1070,1070,1070,1070], font_size=9.2)

    fixed_equity = 96.0 + 325.2 + 9.6 + 2.5
    retained = []
    re = 725.786775 + 30.0 - 48.0
    for n,dv in zip(ni,dividends):
        re = re+n-dv
        retained.append(re)
    other_equity = [206.4544+25.14]*5
    sfp_rows = [
        ["รายการ", "2570", "2571", "2572", "2573", "2574", "ประเภท"],
        ["สินทรัพย์หมุนเวียน", "", "", "", "", "", ""],
        ["เงินสดและรายการเทียบเท่า"] + [money(x) for x in cash] + ["สมมติฐาน"],
        ["ลูกหนี้การค้าและลูกหนี้อื่น"] + [money(x) for x in ar] + ["คำนวณ"],
        ["สินค้าคงเหลือ"] + [money(x) for x in inventory] + ["สมมติฐาน"],
        ["สินทรัพย์ทางการเงินหมุนเวียน"] + [money(x) for x in current_financial] + ["residual × 27.5%"],
        ["สินทรัพย์หมุนเวียนอื่น"] + [money(x) for x in other_current_assets] + ["สมมติฐาน"],
        ["รวมสินทรัพย์หมุนเวียน"] + [money(cash[i]+ar[i]+inventory[i]+current_financial[i]+other_current_assets[i]) for i in range(5)] + ["คำนวณ"],
        ["สินทรัพย์ไม่หมุนเวียน", "", "", "", "", "", ""],
        ["สินทรัพย์ทางการเงินไม่หมุนเวียน"] + [money(x) for x in noncurrent_financial] + ["residual × 72.5%"],
        ["ที่ดิน อาคารและอุปกรณ์"] + [money(x) for x in ppe] + ["Opening + Capex − Dep."],
        ["สินทรัพย์ไม่หมุนเวียนอื่น"] + [money(x) for x in other_noncurrent_assets] + ["สมมติฐาน"],
        ["รวมสินทรัพย์ไม่หมุนเวียน"] + [money(noncurrent_financial[i]+ppe[i]+other_noncurrent_assets[i]) for i in range(5)] + ["คำนวณ"],
        ["รวมสินทรัพย์"] + [money(x) for x in total_assets] + ["คำนวณ"],
        ["หนี้สินหมุนเวียน", "", "", "", "", "", ""],
        ["เจ้าหนี้การค้าและเจ้าหนี้อื่น"] + [money(x) for x in ap] + ["คำนวณ"],
        ["หนี้สินเช่าหมุนเวียน"] + [money(x) for x in current_lease] + ["สมมติฐาน"],
        ["หนี้สินหมุนเวียนอื่น"] + [money(x) for x in other_current_liab] + ["สมมติฐาน"],
        ["รวมหนี้สินหมุนเวียน"] + [money(ap[i]+current_lease[i]+other_current_liab[i]) for i in range(5)] + ["คำนวณ"],
        ["หนี้สินไม่หมุนเวียน", "", "", "", "", "", ""],
        ["หนี้สินเช่าไม่หมุนเวียน"] + [money(x) for x in noncurrent_lease] + ["สมมติฐาน"],
        ["ประมาณการผลประโยชน์พนักงาน"] + [money(x) for x in employee_benefit] + ["สมมติฐาน"],
        ["หนี้สินภาษีเงินได้รอการตัดบัญชี"] + [money(x) for x in deferred_tax] + ["สมมติฐาน"],
        ["รวมหนี้สินไม่หมุนเวียน"] + [money(noncurrent_lease[i]+employee_benefit[i]+deferred_tax[i]) for i in range(5)] + ["คำนวณ"],
        ["รวมหนี้สิน"] + [money(x) for x in total_liab] + ["คำนวณ"],
        ["ส่วนของผู้ถือหุ้น", "", "", "", "", "", ""],
        ["ทุนชำระแล้ว"] + ["96.00"]*5 + ["คงที่"],
        ["ส่วนเกินมูลค่าหุ้นและเงินสำรอง"] + ["337.30"]*5 + ["คงที่"],
        ["กำไรสะสมไม่จัดสรร"] + [money(x) for x in retained] + ["Opening + NI − Dividend"],
        ["องค์ประกอบอื่นของส่วนผู้ถือหุ้น"] + [money(x) for x in other_equity] + ["สมมติฐานคงที่หลัง H1/2569"],
        ["รวมส่วนของผู้ถือหุ้น"] + [money(x) for x in equity] + ["คำนวณ"],
        ["รวมหนี้สินและส่วนของผู้ถือหุ้น"] + [money(total_liab[i]+equity[i]) for i in range(5)] + ["คำนวณ"],
        ["ส่วนต่างตรวจสอบ (สินทรัพย์ − หนี้สิน − ทุน)"] + ["0.00"]*5 + ["ต้องเท่ากับศูนย์"],
    ]
    rebuild_table(sfp, sfp_rows, section_rows={1,8,14,19,25}, widths=[2600,1070,1070,1070,1070,1070,1070],
                  font_size=8.6, numeric_cols={1,2,3,4,5}, emphasis_rows={7,12,13,18,23,24,30,31,32})
    src64 = doc.add_paragraph()
    sfp._tbl.addnext(src64._p)
    set_paragraph_text(src64,
        "ที่มา: ปีฐานจากงบการเงินตรวจสอบแล้ว 2568 [1][2]; ปี 2570–2574 เป็นสมมติฐาน/การคำนวณของผู้จัดทำตามตารางที่ 23 "
        "ฉากทัศน์นี้ไม่ใช่ประมาณการหรือ guidance ของบริษัท", size=10, italic=True, color="666666")
    interpret64 = insert_paragraph_after(doc, src64,
        "นัยเชิงกลยุทธ์ต่อ PG: งบดุลสมดุลทุกปีและ ROE ในตารางที่ 22 ใช้ average equity ที่สร้างจากตารางนี้ "
        "สินทรัพย์ทางการเงินเป็น residual liquidity ภายใต้ cash buffer และอาจเพิ่มขึ้นเมื่อ core profit/ลด inventory มากกว่าการใช้เงิน "
        "จึงไม่ขัดกับการ selective maturity/reallocation ระหว่างปี แต่ยอดจริงต้องแทนด้วย maturity schedule, asset register, dividend policy และ tax model ของบริษัท", size=16)
    set_paragraph_text(p231, "6.5 ข้อจำกัด", style="Heading 2", size=16, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[232],
        "งบแสดงฐานะการเงินที่จัดทำเป็นฉากทัศน์บริหาร ไม่ใช่การพยากรณ์ที่บริษัทรับรอง รายการที่ไม่พบข้อมูลสาธารณะเพียงพอ "
        "ถูกจัดกลุ่มและระบุสูตร/สมมติฐานในตารางที่ 23 ข้อมูลที่ต้องได้จาก P1 และฝ่ายการเงินก่อนใช้ตัดสินใจจริง ได้แก่ margin รายลูกค้า/SKU, "
        "AR/AP ageing, inventory by SKU, capex asset register, depreciation, maturity/yield ของสินทรัพย์ทางการเงิน, ภาษี, เงินปันผล และ committed orders", size=16)
    set_paragraph_text(doc.paragraphs[233],
        "ข้อจำกัดเชิงวิธีการ: น้ำหนัก/Rating/คะแนนคู่แข่งเป็นการประเมินของผู้จัดทำจากข้อมูลสาธารณะ ไม่ใช่ข้อเท็จจริงของบริษัทคู่แข่ง "
        "เมื่อไม่พบข้อมูลสาธารณะเพียงพอ รายงานไม่ใช้ negative-absence claim และระบุว่าไม่พบหลักฐานเพียงพอ "
        "ตัวเลขฉากทัศน์ต้องทบทวนหลัง Gate 0 และทุกไตรมาส; valuation gain/loss ไม่รวมใน forecast เพราะผันผวนและไม่ใช่ผลธุรกิจหลัก", size=16)

    # References: make official trails precise and add current board/shareholder sources.
    set_paragraph_text(doc.paragraphs[241],
        "สำนักงานเศรษฐกิจอุตสาหกรรม กระทรวงอุตสาหกรรม. (2569). สถานการณ์ปี 2568 และแนวโน้มปี 2569 อุตสาหกรรมสิ่งทอและเครื่องนุ่งห่ม. "
        "https://www.oie.go.th/assets/portals/1/fileups/2/files/Industry%20conditions/annual2025trends2026final.pdf [8]", size=13)
    set_paragraph_text(doc.paragraphs[243],
        "European Commission. (2022). EU Strategy for Sustainable and Circular Textiles. "
        "https://environment.ec.europa.eu/strategy/textiles-strategy_en [10]", size=13)
    set_paragraph_text(doc.paragraphs[245],
        "European Union. (2024). Regulation (EU) 2024/3015 on prohibiting products made with forced labour on the Union market. "
        "https://eur-lex.europa.eu/legal-content/EN/ALL/?uri=legissum:5502411 [12]", size=13)
    set_paragraph_text(doc.paragraphs[251],
        "European Commission. (2569). Carbon Border Adjustment Mechanism: CBAM sectors. "
        "https://taxation-customs.ec.europa.eu/carbon-border-adjustment-mechanism/cbam-sectors_en [18]", size=13)
    p254 = doc.paragraphs[254]
    insert_paragraph_before(doc, p254,
        "บริษัท ประชาอาภรณ์ จำกัด (มหาชน). (12 พฤษภาคม 2569). มติแต่งตั้งคณะกรรมการบริษัทและคณะกรรมการชุดย่อย (ทห.013/2569). "
        "https://investor.pg.co.th/wp-content/uploads/2026/05/mati-2-69-E-12-5-69แต่งตั้งกรรมการชุดย่อย.pdf [22]", size=13)
    insert_paragraph_before(doc, p254,
        "บริษัท ประชาอาภรณ์ จำกัด (มหาชน). (13 พฤษภาคม 2569). Shareholders Information. "
        "https://investor.pg.co.th/shareholders-information/ [23]", size=13)

    # Appendix becomes table 25 and explicitly extends evidence-to-result traceability.
    set_paragraph_text(doc.paragraphs[255], "ภาคผนวก Strategic Traceability / Claim-to-Source / Calculation QA", style="Heading 1", size=18, bold=True, color=BLUE)
    set_paragraph_text(doc.paragraphs[256], "ตารางที่ 25 ตารางสอบทานความสอดคล้องและการสืบย้อนกลับของตรรกะเชิงกลยุทธ์", size=14, bold=True)
    trace = doc.tables[-1]
    trace_rows = [
        ["องค์ประกอบในบทหลัง", "สืบย้อนกลับไปยัง", "ตาราง/แหล่งอ้างอิง"],
        ["Current board / shareholder", "มติ 12 พ.ค. 2569 และทะเบียนผู้ถือหุ้น 13 พ.ค. 2569", "[22][23]"],
        ["O1–O4 / T1–T5 และ EFAS", "STEEP, Five Forces, competitor/KSF; SPI ไม่ซ้ำเป็น O", "ตาราง 4–8"],
        ["S1–S6 / W1–W6 และ IFAS", "Resources, VRIO, functions; S6 = SPI/Saha enabler", "ตาราง 9–11"],
        ["SFAS", "คัด/ให้น้ำหนักใหม่จาก EFAS/IFAS", "ตาราง 8,11,12"],
        ["TOWS SO/WO/ST/WT", "ใช้เฉพาะ factors ใน SFAS; รหัสไม่ชน OBJ", "ตาราง 12–13"],
        ["Alternatives A/B/C", "TOWS + competitor arenas + constraints", "ตาราง 6–7,13,16"],
        ["Recommended strategy", "C ชนะ 3.90 และ sensitivity; Gate 0–4", "หัวข้อ 4.5–4.7"],
        ["OBJ1–OBJ9", "Recommended strategy และ scenario targets", "ตาราง 15,22–24"],
        ["P1–P11 / budget", "OBJ/TOWS; owner/timing/budget/KPI/gate", "ตาราง 17–18"],
        ["BSC / risks", "OBJ + projects + early warning/contingency", "ตาราง 19–20"],
        ["Base / Strategic Case", "Actual 2568 + explicit formulas/assumptions", "ตาราง 21–23"],
        ["Projected SFP / ROE", "AR/Inventory/AP/Capex/Dep./Fin assets/Debt/Dividend/RE/Cash", "ตาราง 23–24"],
        ["H1/2569 narrative", "SG&A reduction + valuation reversal; other income down", "ตาราง 3, [4]"],
        ["EU compliance", "CBAM sectors; FLR 14 Dec 2027; cautious DPP", "[10][11][12][18]"],
    ]
    rebuild_table(trace, trace_rows, widths=[2500,4000,2520], font_size=9.5)
    set_paragraph_text(doc.paragraphs[257],
        "การตรวจสอบเชิงตัวเลขและเนื้อหาฉบับปรับปรุงดำเนินการด้วยโปรแกรม: EFAS/IFAS/SFAS/KSF/decision weights รวม 1.00; "
        "weighted score = weight×rating; budget low/high = 102.5/174.0; Gross Profit = Revenue×Gross Margin; core proxy = GP−SG&A; "
        "DIO = Average Inventory/COGS×365; projected assets = liabilities+equity; ROE ใช้ average equity; ไม่พบรหัสวัตถุประสงค์ซ้ำกับ TOWS", size=16)

    # Direct source paragraph below every table.
    sources = [
        "ที่มา: PG One Report/งบการเงิน/คำชี้แจงปี 2568 [1][2][3]; แถวคำนวณระบุในหมายเหตุ",
        "ที่มา: PG One Report ปี 2568 [1]; สัดส่วนเป็นการคำนวณของผู้จัดทำ",
        "ที่มา: คำชี้แจง Q2/2569 ลงวันที่ 11 สิงหาคม 2569 [4]; core proxy เป็นการคำนวณ",
        "ที่มา: PG [1], OIE [8], DITP [9], European Commission/EUR-Lex [10][11][12][18]; การตีความเป็นของผู้จัดทำ",
        "ที่มา: PG [1] และข้อมูลสาธารณะของคู่แข่ง [13][14][15]; ระดับแรงกดดันเป็น analyst assessment",
        "ที่มา: PG [1]; official competitor disclosures [13]–[17][19]; benchmark ไม่ใช่ direct-equivalence claim",
        "ที่มา: [1][14][16]; weights/ratings เป็น analyst estimates และแต่ละสนามมีน้ำหนักรวม 1.00",
        "ที่มา: สังเคราะห์ตาราง 4–7; weights/ratings เป็นการประเมินของผู้จัดทำ (Fact/Calculation/Assumption แยกในเหตุผล)",
        "ที่มา: PG [1][6][7][20][23] และข้อมูลคู่แข่ง [13]–[16]; VRIO เป็น analyst assessment ด้วยสเกล ใช่/บางส่วน/ไม่ใช่",
        "ที่มา: PG [1][4] และการสังเคราะห์ของผู้จัดทำ; negative-absence claims ใช้ถ้อยคำระมัดระวัง",
        "ที่มา: สังเคราะห์จากตาราง 9–10; weights/ratings เป็นการประเมินของผู้จัดทำ",
        "ที่มา: คัดและให้น้ำหนักใหม่จาก EFAS/IFAS; weights/ratings/timing เป็น analyst assessment",
        "ที่มา: สังเคราะห์จาก SFAS; กลยุทธ์เป็นข้อเสนอของผู้จัดทำ",
        "ที่มา: วิสัยทัศน์/พันธกิจบริษัท [1] และผลวิเคราะห์บทที่ 3; ข้อความเสนอเป็นของผู้จัดทำ",
        "ที่มา: สมมติฐานเชิงวางแผนของผู้จัดทำ; เป้าหมายเชื่อมกับตาราง 22–24 ไม่ใช่ company guidance",
        "ที่มา: SFAS/TOWS/competitor evidence; weights/scores/sensitivity คำนวณโดยผู้จัดทำ",
        "ที่มา: แผนของผู้จัดทำ; งบเป็นช่วงสมมติฐานและต้องผ่าน Gate/prerequisite",
        "ที่มา: PG capex ปี 2568 [1] และสมมติฐานของผู้จัดทำ; P8 45–75 ล้านบาท = 15–25 ล้านบาท/ปีโดยเฉลี่ย",
        "ที่มา: OBJ1–OBJ9 และ P1–P11; KPI/targets เป็น management-planning assumptions",
        "ที่มา: Risk assessment ของผู้จัดทำ; early-warning thresholds ต้องกำหนดฐานจริงใน Gate 0/1",
        "ที่มา: ปีฐาน [1][2][3]; ปี 2570–2574 เป็น Base Case assumptions; DIO ใช้ Average Inventory/COGS×365",
        "ที่มา: ปีฐาน [1][2][3]; ปี 2570–2574 เป็น Strategic Case assumptions; ROE ใช้ average equity จากตาราง 24",
        "ที่มา: สมมติฐาน/สูตรของผู้จัดทำ; Management Target Scenario – not company guidance",
        "ที่มา: ปีฐาน [1][2]; ปี 2570–2574 คำนวณจากตาราง 23; Assets = Liabilities + Equity ทุกปี",
        "ที่มา: Strategic traceability audit ของผู้จัดทำ; แหล่งหลัก [1]–[23]",
    ]
    if len(doc.tables) != len(sources):
        raise RuntimeError(f"Expected {len(sources)} tables after insertion, got {len(doc.tables)}")
    for table, src in zip(doc.tables, sources):
        ensure_source_after(doc, table, src)

    # Constrain legacy wide tables to the printable width of their sections.
    # Portrait tables use ~9,100 twips; landscape action/control tables 15,000.
    explicit_widths = {
        3: [1200, 3600, 3600, 700],
        6: [3600, 600, 550, 650, 550, 650, 2500],
        8: [3500, 480, 480, 480, 480, 3680],
        9: [1600, 4800, 2700],
        13: [4100, 5000],
        16: [1800, 2500, 2450, 1800, 900, 1400, 2500, 1650],
        18: [1050, 1900, 3200, 1900, 1900, 1600, 1100, 1350],
    }
    for table_index, widths in explicit_widths.items():
        format_table(doc.tables[table_index], widths=widths, font_size=9.2, header_size=10)

    # Collapse the spillover cover page by removing only unused spacer lines.
    for p in doc.paragraphs[24:30]:
        remove_paragraph(p)

    # The master used a static TOC. Rebuild it compactly so every requested
    # Heading 1/2 section is represented. Page values are refreshed once more
    # from the final Word-rendered PDF during QA.
    toc_entries = [
        ("1. บทสรุปผู้บริหาร", 4, 1),
        ("2. สถานการณ์ปัจจุบัน", 5, 1),
        ("2.1 ประวัติองค์กร", 5, 2),
        ("2.2 โครงสร้างองค์กร", 5, 2),
        ("2.3 คณะกรรมการบริษัทปัจจุบัน", 5, 2),
        ("2.4 วิสัยทัศน์ พันธกิจ วัตถุประสงค์", 6, 2),
        ("2.5 กลยุทธ์ปัจจุบัน", 6, 2),
        ("2.6 ผลประกอบการที่ผ่านมา", 7, 2),
        ("2.7 ผลการดำเนินงานล่าสุด: งวดหกเดือนแรกปี 2569", 9, 2),
        ("3. การวิเคราะห์สภาพแวดล้อมเชิงกลยุทธ์", 9, 1),
        ("3.1 การวิเคราะห์สภาพแวดล้อมภายนอก", 9, 2),
        ("3.2 การวิเคราะห์สภาพแวดล้อมภายใน", 15, 2),
        ("4. การวิเคราะห์ปัจจัยเชิงกลยุทธ์", 18, 1),
        ("4.1 บทสรุปการวิเคราะห์ปัจจัยเชิงกลยุทธ์ (SFAS Matrix)", 18, 2),
        ("4.2 TOWS Matrix", 19, 2),
        ("4.3 การทบทวนวิสัยทัศน์ พันธกิจ และวัตถุประสงค์", 20, 2),
        ("4.4 ทางเลือกเชิงกลยุทธ์ (Strategic Alternatives)", 21, 2),
        ("4.5 เมทริกซ์การตัดสินใจและการทดสอบความไว", 21, 2),
        ("4.6 กลยุทธ์ที่แนะนำ (Recommended Strategy)", 22, 2),
        ("4.7 กลยุทธ์ระดับองค์กร ธุรกิจ และหน้าที่", 23, 2),
        ("5. การนำกลยุทธ์ไปใช้ แผนปฏิบัติการ และการควบคุมประเมินผล", 24, 1),
        ("5.1 แผนการปฏิบัติการ (Action Plan)", 24, 2),
        ("5.2 งบประมาณและฐานการประมาณการ", 27, 2),
        ("5.3 Balanced Scorecard และการควบคุมเชิงกลยุทธ์", 28, 2),
        ("5.4 Risk Register / Early Warning / Contingency", 29, 2),
        ("6. ประมาณการผลการดำเนินงาน", 31, 1),
        ("6.1 กรณีฐาน (Base Case)", 31, 2),
        ("6.2 กรณีกลยุทธ์ (Strategic Case)", 31, 2),
        ("6.3 สมมติฐานของประมาณการ", 32, 2),
        ("6.4 งบแสดงฐานะการเงินประมาณการ", 33, 2),
        ("6.5 ข้อจำกัด", 34, 2),
        ("บรรณานุกรม", 34, 1),
        ("ภาคผนวก Strategic Traceability / Claim-to-Source / Calculation QA", 35, 1),
    ]
    for old in doc.paragraphs[42:65]:
        remove_paragraph(old)
    toc_anchor = doc.paragraphs[65]
    for title, page, level in toc_entries:
        p = toc_anchor.insert_paragraph_before()
        set_toc_line(p, title, page, level)

    # Instruct Word to refresh any remaining fields on open.
    set_update_fields(doc)
    set_core_properties(doc)

    # Ensure all changed/new table text uses the Thai font and no fixed row heights.
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            tr_height = tr_pr.find(qn("w:trHeight"))
            if tr_height is not None:
                tr_pr.remove(tr_height)
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.name is None:
                            set_run_font(r, size=10)

    doc.save(OUTPUT)

    # Console QA summary for the separate audit artifact.
    def check(name, actual, expected, tol=1e-9):
        ok = abs(actual-expected) <= tol
        print(f"{'PASS' if ok else 'FAIL'} | {name} | actual={actual:.6f} expected={expected:.6f}")
        if not ok:
            raise AssertionError(name)

    check("EFAS weight", sum(float(r[1]) for r in efas_rows[2:6] + efas_rows[7:12]), 1.0)
    check("EFAS total", sum(float(r[1])*float(r[2]) for r in efas_rows[2:6] + efas_rows[7:12]), 2.33)
    check("IFAS weight", sum(float(r[1]) for r in ifas_rows[2:8] + ifas_rows[9:15]), 1.0)
    check("IFAS total", sum(float(r[1])*float(r[2]) for r in ifas_rows[2:8] + ifas_rows[9:15]), 2.64)
    check("SFAS weight", sum(float(r[1]) for r in sf_rows[1:15]), 1.0)
    check("SFAS total", sum(float(r[1])*float(r[2]) for r in sf_rows[1:15]), 2.48)
    check("Decision weight", sum([.20,.15,.20,.10,.15,.10,.10]), 1.0)
    check("Decision A", .2*3+.15*4+.2*1+.1*2+.15*1+.1*2+.1*2, 2.15)
    check("Decision B", .2*2+.15*4+.2*4+.1*2+.15*2+.1*2+.1*2, 2.70)
    check("Decision C", .2*5+.15*4+.2*4+.1*4+.15*4+.1*2+.1*3, 3.90)
    for i in range(5):
        check(f"GP {2570+i}", gp[i], rev[i]*gm[i]/100)
        check(f"Core {2570+i}", core[i], gp[i]-sga[i])
        check(f"SFP {2570+i}", total_assets[i], total_liab[i]+equity[i], tol=0.01)
    check("Budget low", 1.5+8+10+6+8+5+45+6+8+5, 102.5)
    check("Budget high", 3+14+18+10+14+9+75+10+13+8, 174.0)
    check("CAGR 2568–2574", (810/605.2)**(1/6)-1, 0.049778522490263155, tol=1e-12)
    print(f"OUTPUT | {OUTPUT}")


if __name__ == "__main__":
    main()
